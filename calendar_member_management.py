import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext
import calendar
from datetime import datetime, date
import json
import os
from typing import Dict, List, Tuple, Optional
import random
from collections import Counter, defaultdict
import copy
import wave
import struct
import math
import tempfile
import threading

# ─────────────── レトロゲーム風サウンドエンジン ───────────────
class SoundManager:
    SAMPLE_RATE = 22050
    def __init__(self):
        self.muted = False
        self._sound_dir = tempfile.mkdtemp(prefix="cal_sfx_")
        self._sounds: Dict[str, str] = {}
        self._pygame_cache: Dict[str, object] = {}
        self._playback_method = self._detect_playback_method()
        self._generate_all_sounds()
        self._preload_sounds()
    def _detect_playback_method(self) -> str:
        try:
            import pygame
            pygame.mixer.init(frequency=self.SAMPLE_RATE, size=-16, channels=1, buffer=256)
            return "pygame"
        except Exception:
            pass
        try:
            import simpleaudio
            return "simpleaudio"
        except ImportError:
            pass
        try:
            import winsound
            return "winsound"
        except ImportError:
            pass
        import shutil
        if shutil.which("aplay"):
            return "aplay"
        if shutil.which("afplay"):
            return "afplay"
        return "none"
    @staticmethod
    def _square_wave(freq, t, duty=0.5):
        if freq == 0: return 0
        phase = (t * freq) % 1.0
        return 1.0 if phase < duty else -1.0
    @staticmethod
    def _triangle_wave(freq, t):
        if freq == 0: return 0
        phase = (t * freq) % 1.0
        return 4.0 * abs(phase - 0.5) - 1.0
    @staticmethod
    def _sine_wave(freq, t):
        if freq == 0: return 0
        return math.sin(2 * math.pi * freq * t)
    @staticmethod
    def _noise(t):
        return random.uniform(-1, 1)
    def _render_wav(self, filename: str, duration: float, generator_func):
        n_samples = int(self.SAMPLE_RATE * duration)
        filepath = os.path.join(self._sound_dir, filename)
        frames = bytearray()
        for i in range(n_samples):
            t = i / self.SAMPLE_RATE
            sample = generator_func(t, duration)
            sample = max(-1.0, min(1.0, sample))
            frames += struct.pack("<h", int(sample * 28000))
        with wave.open(filepath, "w") as wf:
            wf.setnchannels(1)
            wf.setsampwidth(2)
            wf.setframerate(self.SAMPLE_RATE)
            wf.writeframes(bytes(frames))
        return filepath
    def _generate_all_sounds(self):
        self._sounds["bright"]   = self._render_wav("bright.wav",   0.25, self._gen_bright)
        self._sounds["levelup"]  = self._render_wav("levelup.wav",  0.65, self._gen_levelup)
        self._sounds["coin"]     = self._render_wav("coin.wav",     0.35, self._gen_coin)
        self._sounds["snap"]     = self._render_wav("snap.wav",     0.18, self._gen_snap)
        self._sounds["oneup"]    = self._render_wav("oneup.wav",    0.55, self._gen_oneup)
        self._sounds["move"]     = self._render_wav("move.wav",     0.10, self._gen_move)
        self._sounds["vroom"]    = self._render_wav("vroom.wav",    0.50, self._gen_vroom)
        self._sounds["sheen"]    = self._render_wav("sheen.wav",    0.35, self._gen_sheen)
        self._sounds["complete"] = self._render_wav("complete.wav", 1.20, self._gen_complete)
    def _gen_bright(self, t, dur):
        if t < 0.10:
            freq = 659; env = math.cos(t / 0.10 * math.pi * 0.3)
        elif t < 0.23:
            freq = 880; env = 1.0 - ((t - 0.10) / 0.13) ** 0.7
        else:
            return 0
        w = self._sine_wave(freq, t) * 0.55
        w += self._sine_wave(freq * 2, t) * 0.10
        w += self._sine_wave(freq * 3, t) * 0.03
        return w * env * 0.35
    def _gen_levelup(self, t, dur):
        notes = [(0.000,0.100,523),(0.100,0.200,659),(0.200,0.320,784),(0.320,0.620,1047)]
        w = 0
        for start, end, freq in notes:
            if start <= t < end:
                local_t = t - start; note_dur = end - start
                attack = min(1.0, local_t / 0.008)
                decay = max(0, 1.0 - (local_t / note_dur) ** (0.8 if freq==1047 else 0.4))
                env = attack * decay
                w = self._square_wave(freq, t, 0.30) * 0.35 * env
                w += self._triangle_wave(freq, t) * 0.20 * env
                w += self._square_wave(freq * 0.5, t, 0.25) * 0.12 * env
                if freq == 1047:
                    w += self._sine_wave(freq*2, t)*0.08*env + self._sine_wave(freq*3, t)*0.04*env
                break
        return w * 0.55
    def _gen_coin(self, t, dur):
        if t < 0.08:
            freq=988; local_t=t; attack=min(1.0,local_t/0.003); decay=max(0,1.0-(local_t/0.08)**0.6)
            env=attack*decay; w=self._square_wave(freq,t,0.125)*0.45*env+self._sine_wave(freq*2,t)*0.08*env
            return w*0.50
        elif t < 0.30:
            freq=1319; local_t=t-0.08; attack=min(1.0,local_t/0.003); decay=max(0,1.0-(local_t/0.22)**0.8)
            env=attack*decay; w=self._square_wave(freq,t,0.125)*0.40*env+self._sine_wave(freq*2,t)*0.06*env
            return w*0.50
        return 0
    def _gen_snap(self, t, dur):
        if t < 0.02: return self._square_wave(800-t*20000,t,0.5)*0.7
        elif t < 0.08:
            env=1.0-((t-0.02)/0.06); freq=600*(1.0-(t-0.02)/0.06)
            return (self._square_wave(freq,t,0.5)*0.35+self._noise(t)*0.20)*env
        elif t < 0.16: return self._noise(t)*(1.0-((t-0.08)/0.08))*0.10
        return 0
    def _gen_oneup(self, t, dur):
        notes=[(0.000,0.070,659),(0.070,0.140,784),(0.140,0.210,1319),(0.210,0.300,1047),(0.300,0.380,1175),(0.380,0.530,1568)]
        w = 0
        for start, end, freq in notes:
            if start <= t < end:
                local_t=t-start; note_dur=end-start; attack=min(1.0,local_t/0.005)
                decay=max(0,1.0-(local_t/note_dur)**0.5); env=attack*decay
                w=self._square_wave(freq,t,0.25)*0.40*env+self._sine_wave(freq,t)*0.20*env; break
        return w*max(0,1.0-(t/dur)**3)*0.50
    def _gen_move(self, t, dur):
        freq=500+(t/dur)*500; env=max(0,1.0-t/dur)
        return self._square_wave(freq,t,0.25)*env*0.30
    def _gen_vroom(self, t, dur):
        base_freq=80+(t/dur)**1.5*300
        w=self._square_wave(base_freq,t,0.3)*0.30+self._square_wave(base_freq*2,t,0.4)*0.15+self._square_wave(base_freq*3,t,0.5)*0.08
        exhaust=self._noise(t)*0.12; pulse_rate=base_freq/4; pulse=0.7+0.3*math.sin(2*math.pi*pulse_rate*t)
        w=w*pulse+exhaust*pulse
        if t<0.08: env=t/0.08
        elif t<dur-0.10: env=1.0
        else: env=max(0,(dur-t)/0.10)
        return w*env*0.55
    def _gen_sheen(self, t, dur):
        sweep_freq=2000+3000*math.exp(-t*12); attack=t/0.005 if t<0.005 else 1.0
        w=self._sine_wave(sweep_freq,t)*0.25+self._sine_wave(sweep_freq*1.5,t)*0.15+self._sine_wave(sweep_freq*2.3,t)*0.10
        w*=(0.7+0.3*math.sin(2*math.pi*40*t))
        if t<0.03: w+=self._noise(t)*(1.0-t/0.03)*0.30
        if t>0.05:
            ring_env=max(0,1.0-(t-0.05)/0.30)**1.5
            w+=self._sine_wave(3520,t)*ring_env*0.12+self._sine_wave(4186,t)*ring_env*0.08
        return w*attack*max(0,1.0-(t/dur)**1.2)*0.55
    def _gen_complete(self, t, dur):
        notes=[(0.000,0.130,523),(0.130,0.260,659),(0.260,0.400,784),(0.400,0.560,1047),(0.560,0.740,1319),(0.740,0.950,1568),(0.950,1.200,2093)]
        w = 0
        for start, end, freq in notes:
            if start <= t < end:
                local_t=t-start; note_dur=end-start; attack=min(1.0,local_t/0.006)
                decay=max(0,1.0-(local_t/note_dur)**(1.2 if freq==2093 else 0.5)); env=attack*decay
                w=self._square_wave(freq,t,0.25)*0.40*env+self._sine_wave(freq,t)*0.25*env
                w+=self._sine_wave(freq*2,t)*0.12*env+self._sine_wave(freq*3,t)*0.06*env
                if freq==2093:
                    w+=self._sine_wave(freq*1.5,t)*0.08*env+self._sine_wave(freq*2.5,t)*0.04*env
                    w*=(0.85+0.15*math.sin(2*math.pi*8*t))
                break
        if t>1.0: w*=max(0,(dur-t)/0.20)
        return w*0.60
    def _preload_sounds(self):
        if self._playback_method == "pygame":
            try:
                import pygame
                for name, filepath in self._sounds.items():
                    self._pygame_cache[name] = pygame.mixer.Sound(filepath)
            except Exception:
                pass
    def play(self, sound_name: str):
        if self.muted: return
        if self._playback_method == "pygame" and sound_name in self._pygame_cache:
            try: self._pygame_cache[sound_name].play()
            except Exception: pass
            return
        filepath = self._sounds.get(sound_name)
        if not filepath or not os.path.exists(filepath): return
        threading.Thread(target=self._play_sync, args=(filepath,), daemon=True).start()
    def _play_sync(self, filepath: str):
        try:
            if self._playback_method == "simpleaudio":
                import simpleaudio as sa; sa.WaveObject.from_wave_file(filepath).play()
            elif self._playback_method == "winsound":
                import winsound; winsound.PlaySound(filepath, winsound.SND_FILENAME|winsound.SND_ASYNC)
            elif self._playback_method == "aplay":
                import subprocess; subprocess.run(["aplay","-q",filepath],timeout=3,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            elif self._playback_method == "afplay":
                import subprocess; subprocess.run(["afplay",filepath],timeout=3,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        except Exception: pass
    def play_for_label(self, label_name: str):
        if label_name in ("ER","A","B"): self.play("bright")
        elif label_name == "当直": self.play("levelup")
        elif label_name == "外勤": self.play("coin")
        elif label_name == "出張": self.play("snap")
        elif label_name == "休み": self.play("oneup")
    def toggle_mute(self) -> bool:
        self.muted = not self.muted; return self.muted
    def cleanup(self):
        import shutil
        try: shutil.rmtree(self._sound_dir, ignore_errors=True)
        except Exception: pass

CellKey = Tuple[int, int, int]

class CalendarApp:
    def __init__(self, root):
        self.root = root
        self.root.title("カレンダー管理システム（AI支援機能付き）")
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        window_width = min(1600, int(screen_width * 0.9))
        window_height = min(900, int(screen_height * 0.9))
        window_width = max(1200, window_width)
        window_height = max(700, window_height)
        x_position = (screen_width - window_width) // 2
        y_position = (screen_height - window_height) // 2
        self.root.geometry(f"{window_width}x{window_height}+{x_position}+{y_position}")
        if screen_width >= 1600 and screen_height >= 900:
            try: self.root.state("zoomed")
            except: pass
        self.BLOCK_H = 11; self.BLOCK_W = 4
        self.CELL_W = 60; self.CELL_H = 25
        self.LABELS = ["","ER","A","B","当直","明け","外勤","出張","休み"," "," "]
        self.WEEKDAYS = ["月","火","水","木","金","土","日"]
        self.HEADER_H = 30; self.TITLE_H = 45; self.LABEL_COL_W = 50
        self.current_year = 2026; self.current_month = 4
        self.cell_data: Dict[CellKey, str] = {}
        self.cell_colors: Dict[CellKey, str] = {}
        self.cell_sequence: Dict[CellKey, int] = {}
        self.er_marks: set = set()
        self.next_sequence = 1
        self.locked_cells: set = set()
        self.locked_cell_borders: Dict[CellKey, int] = {}
        self.completed_days: set = set()
        self.day_complete_rects: Dict[int, int] = {}
        self.day_memos: Dict[int, str] = {}
        self.memo_windows: Dict[int, tk.Toplevel] = {}
        self.undo_stack: List[Dict] = []
        self.redo_stack: List[Dict] = []
        self.max_history = 30  # ✨ 改善: 3→30に増加
        self.stats_mode = "person"
        self.value_colors: Dict[str, str] = {
            "ER":"#98fb98","A":"#ffcc99","B":"#add8e6","当直":"#ffb6c1",
            "外勤":"#dda0dd","出張":"#f0e68c","休み":"#e0ffff",
            "秋":"#FFE4B2","桜":"#FFD1DC","小":"#D8D8E8","金":"#C5EDD6",
            "坪":"#FFDCB5","東":"#DDD0F5","長":"#BDE3F8","矢":"#FFF0B5","宮":"#E0E0E0",
        }
        self._rect_ids: Dict[CellKey, int] = {}
        self._text_ids: Dict[CellKey, int] = {}
        self._cell_bounds: Dict[CellKey, Tuple[int,int,int,int]] = {}
        self._edit_entry: Optional[tk.Entry] = None
        self._edit_key: Optional[CellKey] = None
        self._edit_window_id: Optional[int] = None
        self._drag_source: Optional[CellKey] = None
        self._drag_start_cx: float = 0
        self._drag_start_cy: float = 0
        self._dragging: bool = False
        self._drag_ghost_rect: Optional[int] = None
        self._drag_ghost_text: Optional[int] = None
        self._drag_highlight: Optional[int] = None
        self.DRAG_THRESHOLD = 5
        self.selected_value: Optional[str] = None
        self.mask_opacity = 0.6
        self.special_borders: Dict[CellKey, str] = {}
        self.special_border_ids: Dict[CellKey, List[int]] = {}
        self._day_positions: Dict[int, Tuple[int,int,int,int]] = {}
        self.sound_manager = SoundManager()
        self.team_rules = ""
        self.members_data = []
        self._next_member_id = 1
        self._load_default_members()
        self.sub_cell_data: Dict[Tuple[str,int,str], str] = {}
        self._sub_cell_bounds: Dict[Tuple[str,int,str], Tuple[int,int,int,int]] = {}
        self._sub_edit_entry = None
        self._sub_edit_key: Optional[Tuple] = None
        self._sub_edit_win_id = None
        self._sub_ctx_menu: Optional[tk.Menu] = None
        self.special_items: Dict[str, List[str]] = {"外勤":[],"委員会":[],"コース":[],"訓練":[]}
        self._special_item_windows: Dict[str, tk.Toplevel] = {}
        self._hp_tooltip_win: Optional[tk.Toplevel] = None
        self._hp_tooltip_areas: List[tuple] = []
        self.ai_conversation_history: List[Dict] = []
        self.ai_latest_proposal: Optional[Dict] = None
        self.ai_preview_active: bool = False
        self.ai_preview_data: Dict = {}
        self._keyring_service = "CalendarAI_Anthropic"
        self._keyring_username = "api_key"
        self._keyring_available = self._check_keyring()
        self.api_key: str = self._load_api_key()
        # ✨ 改善: 祝日セット（カレンダー生成時に設定）
        self._current_holidays: set = set()
        # ✨ 改善: 自動バックアップ設定
        self._autosave_interval = 5 * 60 * 1000  # 5分
        self.setup_ui()
        self.root.update_idletasks()
        self.root.after(100, self.create_calendar)
        # ✨ 改善: 自動バックアップ開始
        self.root.after(self._autosave_interval, self._autosave)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    # ✨ 改善: 日本の祝日を返すメソッド
    def _get_japanese_holidays(self, year: int, month: int) -> set:
        """日本の主要祝日を返す（固定祝日ベース）"""
        fixed = {
            (1,1):"元日",(2,11):"建国記念の日",(2,23):"天皇誕生日",
            (4,29):"昭和の日",(5,3):"憲法記念日",(5,4):"みどりの日",(5,5):"こどもの日",
            (8,11):"山の日",(11,3):"文化の日",(11,23):"勤労感謝の日",
        }
        # 春分・秋分の近似
        spring = 20 if year <= 2099 else 21
        autumn = 23 if year <= 2044 else 22
        fixed[(3, spring)] = "春分の日"
        fixed[(9, autumn)] = "秋分の日"
        result = set()
        for (m, d), name in fixed.items():
            if m == month:
                result.add(d)
        return result

    # ✨ 改善: 自動バックアップ
    def _autosave(self):
        """5分ごとの自動バックアップ保存"""
        if self.cell_data:
            backup_dir = os.path.join(tempfile.gettempdir(), "calendar_autosave")
            os.makedirs(backup_dir, exist_ok=True)
            filename = os.path.join(backup_dir, f"autosave_{self.current_year}_{self.current_month}.json")
            try:
                data = {
                    "year": self.current_year, "month": self.current_month,
                    "cell_data": {f"{d},{i},{j}": v for (d,i,j),v in self.cell_data.items()},
                    "cell_colors": {f"{d},{i},{j}": v for (d,i,j),v in self.cell_colors.items()},
                    "cell_sequence": {f"{d},{i},{j}": v for (d,i,j),v in self.cell_sequence.items()},
                    "value_colors": self.value_colors, "next_sequence": self.next_sequence,
                    "members_data": self.members_data, "team_rules": self.team_rules,
                }
                with open(filename, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False)
            except Exception:
                pass
        self.root.after(self._autosave_interval, self._autosave)

    # ✨ 改善: 前月・次月・今月ナビゲーション
    def _prev_month(self):
        y, m = int(self.year_var.get()), int(self.month_var.get())
        if m == 1: y -= 1; m = 12
        else: m -= 1
        self.year_var.set(str(y)); self.month_var.set(str(m))
        self.create_calendar()

    def _next_month(self):
        y, m = int(self.year_var.get()), int(self.month_var.get())
        if m == 12: y += 1; m = 1
        else: m += 1
        self.year_var.set(str(y)); self.month_var.set(str(m))
        self.create_calendar()

    def _goto_today(self):
        today = date.today()
        self.year_var.set(str(today.year)); self.month_var.set(str(today.month))
        self.create_calendar()

    # ✨ 改善: メンバーの色をセルに反映するメソッド（バグ修正）
    def _apply_member_color_to_cells(self, display_name: str, color: str, old_display_name: str = None):
        """メンバーの色をvalue_colorsに登録し、既存セルの色を更新する"""
        self.value_colors[display_name] = color
        if old_display_name and old_display_name != display_name:
            # 表示名が変わった場合: セルデータを新名前に更新
            keys_to_update = [k for k, v in self.cell_data.items() if v == old_display_name]
            for k in keys_to_update:
                self.cell_data[k] = display_name
                self.cell_colors[k] = color
            # サブカレンダーデータも更新
            sub_keys = [k for k in list(self.sub_cell_data.keys()) if k[0] == old_display_name]
            for k in sub_keys:
                new_k = (display_name, k[1], k[2])
                self.sub_cell_data[new_k] = self.sub_cell_data.pop(k)
            if old_display_name in self.value_colors:
                del self.value_colors[old_display_name]
        else:
            # 同じ表示名: 色だけ更新
            for k, v in self.cell_data.items():
                if v == display_name:
                    self.cell_colors[k] = color
        self._redraw_all_cells()
        self.draw_statistics()

    def setup_ui(self):
        self.root.minsize(1200, 700)
        screen_width = self.root.winfo_screenwidth()
        is_small_screen = screen_width < 1600
        if is_small_screen:
            btn_kw = {"font":("Meiryo UI",10),"width":12,"height":1}
            undo_redo_kw = {"font":("Meiryo UI",10),"width":7,"height":1}
            ai_btn_kw = {"font":("Meiryo UI",10,"bold"),"width":10,"height":1}
            sound_btn_kw = {"font":("Meiryo UI",9,"bold"),"width":8,"height":1}
            nav_kw = {"font":("Meiryo UI",10,"bold"),"width":6,"height":1}
        else:
            btn_kw = {"font":("Meiryo UI",11),"width":15,"height":1}
            undo_redo_kw = {"font":("Meiryo UI",11),"width":8,"height":1}
            ai_btn_kw = {"font":("Meiryo UI",11,"bold"),"width":12,"height":1}
            sound_btn_kw = {"font":("Meiryo UI",10,"bold"),"width":10,"height":1}
            nav_kw = {"font":("Meiryo UI",11,"bold"),"width":7,"height":1}
        top_frame = tk.Frame(self.root, bg="#f0f0f0")
        top_frame.pack(fill=tk.X, padx=10, pady=(10 if is_small_screen else 6))
        if is_small_screen:
            row1 = tk.Frame(top_frame, bg="#f0f0f0")
            row1.pack(fill=tk.X, pady=2)
            tk.Label(row1, text="年:", bg="#f0f0f0", font=("Meiryo UI",11)).pack(side=tk.LEFT, padx=5)
            self.year_var = tk.StringVar(value=str(self.current_year))
            ttk.Spinbox(row1, from_=2020, to=2030, textvariable=self.year_var, width=8, font=("Meiryo UI",10)).pack(side=tk.LEFT, padx=5)
            tk.Label(row1, text="月:", bg="#f0f0f0", font=("Meiryo UI",11)).pack(side=tk.LEFT, padx=5)
            self.month_var = tk.StringVar(value=str(self.current_month))
            ttk.Spinbox(row1, from_=1, to=12, textvariable=self.month_var, width=8, font=("Meiryo UI",10)).pack(side=tk.LEFT, padx=5)
            # ✨ 改善: 前月・今月・次月ボタン
            tk.Button(row1, text="◀", command=self._prev_month, bg="#607d8b", fg="white", **nav_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="今月", command=self._goto_today, bg="#607d8b", fg="white", **nav_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="▶", command=self._next_month, bg="#607d8b", fg="white", **nav_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="カレンダー生成", command=self.create_calendar, bg="#2d5a7b", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row1, text="◀ 戻る", command=self.undo, bg="#6c757d", fg="white", **undo_redo_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="進む ▶", command=self.redo, bg="#6c757d", fg="white", **undo_redo_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="クリア", command=self.clear_data, bg="#e76f51", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            self.sound_btn = tk.Button(row1, text="🔊 SE:ON", command=self._toggle_sound, bg="#4CAF50", fg="white", cursor="hand2", **sound_btn_kw)
            self.sound_btn.pack(side=tk.RIGHT, padx=5)
            self.ai_open_btn = tk.Button(row1, text="🤖 AI支援", command=self._open_ai_window, bg="#805ad5", fg="white", cursor="hand2", **ai_btn_kw)
            self.ai_open_btn.pack(side=tk.RIGHT, padx=5)
            row2 = tk.Frame(top_frame, bg="#f0f0f0")
            row2.pack(fill=tk.X, pady=2)
            tk.Button(row2, text="Excel出力", command=self.export_to_excel, bg="#f4a261", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row2, text="データ保存", command=self.save_data, bg="#264653", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row2, text="データ読込", command=self.load_data, bg="#2a9d8f", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row2, text="👥 メンバー管理", command=self.open_member_management, bg="#8b5cf6", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            row3 = tk.Frame(top_frame, bg="#f0f0f0")
            row3.pack(fill=tk.X, pady=2)
            special_btn_kw = {"font":("Meiryo UI",10,"bold"),"width":10,"height":1}
            tk.Label(row3, text="特別項目:", bg="#f0f0f0", font=("Meiryo UI",10)).pack(side=tk.LEFT, padx=(5,2))
            tk.Button(row3, text="🚗 外勤", command=lambda: self._open_special_item_window("外勤"), bg="#5b8dd9", fg="white", cursor="hand2", **special_btn_kw).pack(side=tk.LEFT, padx=3)
            tk.Button(row3, text="🏛 委員会", command=lambda: self._open_special_item_window("委員会"), bg="#5b8dd9", fg="white", cursor="hand2", **special_btn_kw).pack(side=tk.LEFT, padx=3)
            tk.Button(row3, text="📚 コース", command=lambda: self._open_special_item_window("コース"), bg="#5b8dd9", fg="white", cursor="hand2", **special_btn_kw).pack(side=tk.LEFT, padx=3)
            tk.Button(row3, text="🎯 訓練", command=lambda: self._open_special_item_window("訓練"), bg="#5b8dd9", fg="white", cursor="hand2", **special_btn_kw).pack(side=tk.LEFT, padx=3)
        else:
            row1 = tk.Frame(top_frame, bg="#f0f0f0")
            row1.pack(fill=tk.X, pady=2)
            tk.Label(row1, text="年:", bg="#f0f0f0", font=("Meiryo UI",12)).pack(side=tk.LEFT, padx=5)
            self.year_var = tk.StringVar(value=str(self.current_year))
            ttk.Spinbox(row1, from_=2020, to=2030, textvariable=self.year_var, width=10, font=("Meiryo UI",11)).pack(side=tk.LEFT, padx=5)
            tk.Label(row1, text="月:", bg="#f0f0f0", font=("Meiryo UI",12)).pack(side=tk.LEFT, padx=5)
            self.month_var = tk.StringVar(value=str(self.current_month))
            ttk.Spinbox(row1, from_=1, to=12, textvariable=self.month_var, width=10, font=("Meiryo UI",11)).pack(side=tk.LEFT, padx=5)
            # ✨ 改善: 前月・今月・次月ボタン
            tk.Button(row1, text="◀ 前月", command=self._prev_month, bg="#607d8b", fg="white", **nav_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="今月", command=self._goto_today, bg="#607d8b", fg="white", **nav_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="次月 ▶", command=self._next_month, bg="#607d8b", fg="white", **nav_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="カレンダー生成", command=self.create_calendar, bg="#2d5a7b", fg="white", **btn_kw).pack(side=tk.LEFT, padx=10)
            tk.Button(row1, text="◀ 戻る", command=self.undo, bg="#6c757d", fg="white", **undo_redo_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="進む ▶", command=self.redo, bg="#6c757d", fg="white", **undo_redo_kw).pack(side=tk.LEFT, padx=2)
            tk.Button(row1, text="クリア", command=self.clear_data, bg="#e76f51", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row1, text="Excel出力", command=self.export_to_excel, bg="#f4a261", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row1, text="データ保存", command=self.save_data, bg="#264653", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row1, text="データ読込", command=self.load_data, bg="#2a9d8f", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            tk.Button(row1, text="👥 メンバー管理", command=self.open_member_management, bg="#8b5cf6", fg="white", **btn_kw).pack(side=tk.LEFT, padx=5)
            self.sound_btn = tk.Button(row1, text="🔊 SE:ON", command=self._toggle_sound, bg="#4CAF50", fg="white", cursor="hand2", **sound_btn_kw)
            self.sound_btn.pack(side=tk.RIGHT, padx=10)
            self.ai_open_btn = tk.Button(row1, text="🤖 AI支援", command=self._open_ai_window, bg="#805ad5", fg="white", cursor="hand2", **ai_btn_kw)
            self.ai_open_btn.pack(side=tk.RIGHT, padx=5)
            row2 = tk.Frame(top_frame, bg="#f0f0f0")
            row2.pack(fill=tk.X, pady=2)
            special_btn_kw = {"font":("Meiryo UI",10,"bold"),"width":10,"height":1}
            tk.Label(row2, text="特別項目:", bg="#f0f0f0", font=("Meiryo UI",10)).pack(side=tk.LEFT, padx=(5,2))
            for label, cat in [("🚗 外勤","外勤"),("🏛 委員会","委員会"),("📚 コース","コース"),("🎯 訓練","訓練")]:
                tk.Button(row2, text=label, command=lambda c=cat: self._open_special_item_window(c), bg="#5b8dd9", fg="white", cursor="hand2", **special_btn_kw).pack(side=tk.LEFT, padx=3)
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.tab_main = tk.Frame(self.notebook)
        self.notebook.add(self.tab_main, text="📅 メインカレンダー")
        main_sync_bar = tk.Frame(self.tab_main, bg="#e8f4f8", bd=1, relief=tk.GROOVE)
        main_sync_bar.pack(fill=tk.X, padx=5, pady=(5,0))
        tk.Button(main_sync_bar, text="🔄 カレンダー同期（メインをサブに合わせる）", command=self._on_main_sync_btn, bg="#2a9d8f", fg="white", font=("Meiryo UI",10,"bold"), relief=tk.FLAT, padx=15, pady=3, cursor="hand2").pack(side=tk.LEFT, padx=8, pady=4)
        tk.Label(main_sync_bar, text="← サブカレンダーの内容でメインを上書き", bg="#e8f4f8", font=("Meiryo UI",9), fg="#444").pack(side=tk.LEFT, padx=4)
        self.paned = tk.PanedWindow(self.tab_main, orient=tk.HORIZONTAL, sashwidth=5, bg="#d9d9d9")
        self.paned.pack(fill=tk.BOTH, expand=True)
        self.left_frame = tk.Frame(self.paned, bg="white")
        self.paned.add(self.left_frame, minsize=800, stretch="always")
        v_scroll = tk.Scrollbar(self.left_frame, orient=tk.VERTICAL)
        v_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        h_scroll = tk.Scrollbar(self.left_frame, orient=tk.HORIZONTAL)
        h_scroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.canvas = tk.Canvas(self.left_frame, bg="white", yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        v_scroll.config(command=self.canvas.yview)
        h_scroll.config(command=self.canvas.xview)
        self.canvas.bind("<Double-1>", self._on_double_click)
        self.canvas.bind("<Button-1>", self._on_click)
        self.canvas.bind("<B1-Motion>", self._on_motion)
        self.canvas.bind("<ButtonRelease-1>", self._on_release)
        self.canvas.bind("<Button-3>", self._on_right_click)
        self._bind_mousewheel(self.canvas, self.canvas)
        self.root.bind("<Control-z>", lambda e: self.undo())
        self.root.bind("<Control-y>", lambda e: self.redo())
        self.right_frame = tk.Frame(self.paned, bg="#f9f9f9", width=400)
        self.paned.add(self.right_frame, minsize=350, stretch="never")
        stats_header = tk.Frame(self.right_frame, bg="#f9f9f9")
        stats_header.pack(fill=tk.X, pady=5)
        tk.Label(stats_header, text="【集計・グラフ】", font=("Meiryo UI",13,"bold"), bg="#f9f9f9").pack(side=tk.TOP, pady=5)
        button_frame = tk.Frame(stats_header, bg="#f9f9f9")
        button_frame.pack(side=tk.TOP, pady=5)
        self.btn_person = tk.Button(button_frame, text="人別集計", command=lambda: self.switch_stats_mode("person"), bg="#4a90e2", fg="white", font=("Meiryo UI",10,"bold"), width=12, relief=tk.SUNKEN)
        self.btn_person.pack(side=tk.LEFT, padx=5)
        self.btn_category = tk.Button(button_frame, text="項目別集計", command=lambda: self.switch_stats_mode("category"), bg="#7c8a9e", fg="white", font=("Meiryo UI",10), width=12, relief=tk.RAISED)
        self.btn_category.pack(side=tk.LEFT, padx=5)
        stats_container = tk.Frame(self.right_frame, bg="white")
        stats_container.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        stats_vscroll = tk.Scrollbar(stats_container, orient=tk.VERTICAL)
        stats_vscroll.pack(side=tk.RIGHT, fill=tk.Y)
        stats_hscroll = tk.Scrollbar(stats_container, orient=tk.HORIZONTAL)
        stats_hscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.stats_canvas = tk.Canvas(stats_container, bg="white", highlightthickness=0, yscrollcommand=stats_vscroll.set, xscrollcommand=stats_hscroll.set)
        self.stats_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        stats_vscroll.config(command=self.stats_canvas.yview)
        stats_hscroll.config(command=self.stats_canvas.xview)
        self._bind_mousewheel(stats_container, self.stats_canvas)
        self._bind_mousewheel(self.stats_canvas, self.stats_canvas)
        self.stats_canvas.bind("<Motion>", self._on_stats_motion)
        self.stats_canvas.bind("<Leave>", self._hide_hp_tooltip)
        self.tab_sub = tk.Frame(self.notebook)
        self.notebook.add(self.tab_sub, text="📋 サブカレンダー")
        sub_sync_bar = tk.Frame(self.tab_sub, bg="#f0f8ee", bd=1, relief=tk.GROOVE)
        sub_sync_bar.pack(fill=tk.X, padx=5, pady=(5,0))
        tk.Button(sub_sync_bar, text="🔄 カレンダー同期（サブをメインに合わせる）", command=self._on_sub_sync_btn, bg="#2d6a4f", fg="white", font=("Meiryo UI",10,"bold"), relief=tk.FLAT, padx=15, pady=3, cursor="hand2").pack(side=tk.LEFT, padx=8, pady=4)
        tk.Label(sub_sync_bar, text="← メインカレンダーの内容でサブを上書き", bg="#f0f8ee", font=("Meiryo UI",9), fg="#444").pack(side=tk.LEFT, padx=4)
        sub_vscroll = tk.Scrollbar(self.tab_sub, orient=tk.VERTICAL)
        sub_vscroll.pack(side=tk.RIGHT, fill=tk.Y)
        sub_hscroll = tk.Scrollbar(self.tab_sub, orient=tk.HORIZONTAL)
        sub_hscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.sub_canvas = tk.Canvas(self.tab_sub, bg="white", yscrollcommand=sub_vscroll.set, xscrollcommand=sub_hscroll.set)
        self.sub_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sub_vscroll.config(command=self.sub_canvas.yview)
        sub_hscroll.config(command=self.sub_canvas.xview)
        self._bind_mousewheel(self.sub_canvas, self.sub_canvas)
        self.sub_canvas.bind("<Button-1>", self._on_sub_canvas_click)
        self.sub_canvas.bind("<Button-3>", self._on_sub_canvas_right_click)
        self.tab_personal = tk.Frame(self.notebook)
        self.notebook.add(self.tab_personal, text="👤 個人別一覧")
        personal_vscroll = tk.Scrollbar(self.tab_personal, orient=tk.VERTICAL)
        personal_vscroll.pack(side=tk.RIGHT, fill=tk.Y)
        personal_hscroll = tk.Scrollbar(self.tab_personal, orient=tk.HORIZONTAL)
        personal_hscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.personal_canvas = tk.Canvas(self.tab_personal, bg="white", yscrollcommand=personal_vscroll.set, xscrollcommand=personal_hscroll.set)
        self.personal_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        personal_vscroll.config(command=self.personal_canvas.yview)
        personal_hscroll.config(command=self.personal_canvas.xview)
        self.notebook.bind("<<NotebookTabChanged>>", self._on_tab_changed)
        self.ai_window = None

    def _bind_mousewheel(self, widget, target_canvas):
        def _on_wheel(event):
            if event.delta:
                target_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
            elif event.num == 4:
                target_canvas.yview_scroll(-3, "units")
            elif event.num == 5:
                target_canvas.yview_scroll(3, "units")
        widget.bind("<MouseWheel>", _on_wheel)
        widget.bind("<Button-4>", _on_wheel)
        widget.bind("<Button-5>", _on_wheel)

    def _check_keyring(self) -> bool:
        try:
            import keyring
            keyring.get_credential(self._keyring_service, self._keyring_username)
            return True
        except ImportError:
            return False
        except Exception:
            return False

    def _load_api_key(self) -> str:
        if self._keyring_available:
            try:
                import keyring
                key = keyring.get_password(self._keyring_service, self._keyring_username)
                if key:
                    return key
            except Exception:
                pass
        return ""

    def _save_api_key(self, key: str):
        if self._keyring_available:
            try:
                import keyring
                keyring.set_password(self._keyring_service, self._keyring_username, key.strip())
                return True
            except Exception:
                pass
        return False

    def _delete_api_key(self):
        if self._keyring_available:
            try:
                import keyring
                keyring.delete_password(self._keyring_service, self._keyring_username)
            except Exception:
                pass

    def _prompt_api_key(self):
        parent = self.ai_window if self.ai_window else self.root
        dialog = tk.Toplevel(parent)
        dialog.title("🔑 APIキー設定")
        dialog.geometry("560x310")
        dialog.resizable(False, False)
        dialog.configure(bg="#f5f7fa")
        dialog.transient(parent)
        dialog.grab_set()
        tk.Label(dialog, text="🔑 Anthropic APIキーを入力してください",
                 font=("Meiryo UI", 12, "bold"), bg="#f5f7fa").pack(pady=(20, 5))
        tk.Label(dialog, text="https://console.anthropic.com/ から取得できます",
                 font=("Meiryo UI", 9), bg="#f5f7fa", fg="#718096").pack(pady=(0, 10))
        key_frame = tk.Frame(dialog, bg="#f5f7fa")
        key_frame.pack(fill=tk.X, padx=30, pady=5)
        key_var = tk.StringVar(value=self.api_key)
        key_entry = tk.Entry(key_frame, textvariable=key_var, show="•",
                             font=("Consolas", 11), width=48)
        key_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        entry_menu = tk.Menu(key_entry, tearoff=0)
        entry_menu.add_command(label="切り取り (Ctrl+X)", command=lambda: key_entry.event_generate("<<Cut>>"))
        entry_menu.add_command(label="コピー (Ctrl+C)", command=lambda: key_entry.event_generate("<<Copy>>"))
        entry_menu.add_command(label="貼り付け (Ctrl+V)", command=lambda: key_entry.event_generate("<<Paste>>"))
        entry_menu.add_separator()
        entry_menu.add_command(label="全選択 (Ctrl+A)", command=lambda: (key_entry.select_range(0, tk.END), key_entry.icursor(tk.END)))
        key_entry.bind("<Button-3>", lambda e: entry_menu.tk_popup(e.x_root, e.y_root))
        show_var = tk.BooleanVar(value=False)
        def toggle_show():
            key_entry.config(show="" if show_var.get() else "•")
        tk.Checkbutton(key_frame, text="表示", variable=show_var, command=toggle_show,
                       bg="#f5f7fa", font=("Meiryo UI", 9)).pack(side=tk.RIGHT, padx=5)
        save_var = tk.BooleanVar(value=self._keyring_available)
        save_cb = tk.Checkbutton(dialog, text="次回以降のためにキーを保存する（OS資格情報ストア使用）",
                       variable=save_var, bg="#f5f7fa", font=("Meiryo UI", 9))
        save_cb.pack(pady=5)
        if self._keyring_available:
            storage_info = "🔒 Windows資格情報マネージャー / macOS Keychain / Linux Secret Service に暗号化保存されます"
            info_color = "#38a169"
        else:
            storage_info = ("⚠️ keyringライブラリ未導入のため保存できません（今回の起動中のみ有効）\n"
                           "　　pip install keyring でインストールすると安全に保存できます")
            info_color = "#e53e3e"
            save_cb.config(state=tk.DISABLED)
            save_var.set(False)
        tk.Label(dialog, text=storage_info, font=("Meiryo UI", 8),
                 bg="#f5f7fa", fg=info_color, wraplength=480, justify=tk.LEFT).pack(padx=30, pady=(0, 10))
        btn_frame = tk.Frame(dialog, bg="#f5f7fa")
        btn_frame.pack(pady=10)
        result = {"ok": False}
        def on_ok():
            k = key_var.get().strip()
            if not k:
                messagebox.showwarning("警告", "APIキーを入力してください", parent=dialog)
                return
            self.api_key = k
            if save_var.get():
                saved = self._save_api_key(k)
                if not saved:
                    messagebox.showwarning("警告", "キーの保存に失敗しました。今回の起動中のみ有効です。", parent=dialog)
            result["ok"] = True
            dialog.destroy()
        def on_cancel():
            dialog.destroy()
        tk.Button(btn_frame, text="✅ 設定", command=on_ok,
                  bg="#48bb78", fg="white", font=("Meiryo UI", 10, "bold"),
                  relief=tk.FLAT, padx=20, pady=6, cursor="hand2").pack(side=tk.LEFT, padx=10)
        tk.Button(btn_frame, text="キャンセル", command=on_cancel,
                  bg="#a0aec0", fg="white", font=("Meiryo UI", 10),
                  relief=tk.FLAT, padx=15, pady=6, cursor="hand2").pack(side=tk.LEFT, padx=10)
        key_entry.focus_force()
        key_entry.bind("<Return>", lambda e: on_ok())
        dialog.wait_window()
        return result["ok"]

    def _open_ai_window(self):
        if self.ai_window is not None:
            try:
                self.ai_window.lift()
                self.ai_window.focus_force()
                if hasattr(self, 'ai_input') and self.ai_input.winfo_exists():
                    self.ai_input.focus_force()
                return
            except tk.TclError:
                self.ai_window = None
        win = tk.Toplevel(self.root)
        win.title("🤖 Claude AI 勤務表支援")
        try:
            self.root.update_idletasks()
            screen_width = self.root.winfo_screenwidth()
            screen_height = self.root.winfo_screenheight()
            ai_width = 400
            max_height = int(screen_height * 0.85)
            ai_height = min(700, max_height)
            ai_height = max(ai_height, 500)
            ai_x = screen_width - ai_width - 10
            ai_y = 30
            if ai_y + ai_height > screen_height:
                ai_y = max(0, screen_height - ai_height - 40)
            win.geometry(f"{ai_width}x{ai_height}+{ai_x}+{ai_y}")
        except:
            win.geometry("400x600")
        win.minsize(350, 500)
        win.configure(bg="#f5f7fa")
        self.ai_window = win
        win.protocol("WM_DELETE_WINDOW", self._close_ai_window)
        header = tk.Frame(win, bg="#4a5568", height=50)
        header.pack(fill=tk.X, side=tk.TOP)
        header.pack_propagate(False)
        tk.Label(header, text="🤖 Claude AI 勤務表支援", font=("Meiryo UI", 14, "bold"),
                 bg="#4a5568", fg="white").pack(side=tk.LEFT, padx=15, pady=10)
        tk.Button(header, text="🔄 リセット", command=self._reset_ai_conversation,
                  bg="#e53e3e", fg="white", font=("Meiryo UI", 9), relief=tk.FLAT, padx=10, pady=3).pack(side=tk.RIGHT, padx=15)
        tk.Button(header, text="🔑 APIキー", command=self._prompt_api_key,
                  bg="#718096", fg="white", font=("Meiryo UI", 9), relief=tk.FLAT, padx=10, pady=3).pack(side=tk.RIGHT, padx=0)
        main_container = tk.Frame(win, bg="#f5f7fa")
        main_container.pack(fill=tk.BOTH, expand=True, padx=12, pady=12)
        chat_frame = tk.Frame(main_container, bg="white", relief=tk.RIDGE, bd=2)
        chat_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
        chat_scroll = tk.Scrollbar(chat_frame)
        chat_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.ai_chat_display = tk.Text(chat_frame, wrap=tk.WORD, font=("Meiryo UI", 10),
                                       bg="white", fg="#2d3748", yscrollcommand=chat_scroll.set,
                                       state=tk.DISABLED, padx=12, pady=12)
        self.ai_chat_display.pack(fill=tk.BOTH, expand=True)
        chat_scroll.config(command=self.ai_chat_display.yview)
        self.ai_chat_display.tag_config("user", foreground="#2b6cb0", font=("Meiryo UI", 10, "bold"))
        self.ai_chat_display.tag_config("assistant", foreground="#38a169", font=("Meiryo UI", 10, "bold"))
        self.ai_chat_display.tag_config("system", foreground="#718096", font=("Meiryo UI", 9, "italic"))
        self.ai_chat_display.tag_config("proposal", background="#fef5e7", relief=tk.RAISED, borderwidth=1)
        examples_frame = tk.Frame(main_container, bg="#f5f7fa")
        examples_frame.pack(fill=tk.X, pady=(0, 6))
        tk.Label(examples_frame, text="💡 例:", bg="#f5f7fa", font=("Meiryo UI", 9)).pack(side=tk.LEFT, padx=(0, 4))
        for ex in ["休みを均等に配分して", "桜と小の当直を被らないように", "金のER勤務を増やして"]:
            tk.Button(examples_frame, text=ex, command=lambda e=ex: self._insert_example(e),
                      bg="#e2e8f0", fg="#2d3748", font=("Meiryo UI", 8),
                      relief=tk.FLAT, padx=6, pady=2, cursor="hand2").pack(side=tk.LEFT, padx=2)
        input_frame = tk.Frame(main_container, bg="white", relief=tk.RIDGE, bd=2)
        input_frame.pack(fill=tk.X)
        self.ai_input = tk.Text(input_frame, height=3, wrap=tk.WORD, font=("Meiryo UI", 11),
                                bg="#fffff8", fg="#2d3748", relief=tk.FLAT, padx=10, pady=8,
                                insertbackground="black", insertwidth=2)
        self.ai_input.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(4, 0), pady=4)
        self.ai_input.bind("<Control-Return>", lambda e: (self._send_ai_message(), "break")[-1])
        self.ai_input.bind("<Return>", self._on_ai_enter_key)
        btn_frame = tk.Frame(input_frame, bg="white")
        btn_frame.pack(side=tk.RIGHT, padx=8, pady=8)
        self.ai_send_btn = tk.Button(btn_frame, text="📤 送信\n(Enter)", command=self._send_ai_message,
                                     bg="#4299e1", fg="white", font=("Meiryo UI", 10, "bold"),
                                     relief=tk.FLAT, padx=12, pady=8, cursor="hand2")
        self.ai_send_btn.pack(fill=tk.BOTH, expand=True)
        self.ai_action_frame = tk.Frame(main_container, bg="#f5f7fa")
        self.ai_preview_btn = tk.Button(self.ai_action_frame, text="👁️ プレビュー",
                                        command=self._preview_ai_proposal, bg="#805ad5", fg="white",
                                        font=("Meiryo UI", 10, "bold"), relief=tk.FLAT, padx=15, pady=8, cursor="hand2")
        self.ai_preview_btn.pack(side=tk.LEFT, padx=4)
        self.ai_apply_btn = tk.Button(self.ai_action_frame, text="✅ 適用",
                                      command=self._apply_ai_proposal, bg="#48bb78", fg="white",
                                      font=("Meiryo UI", 10, "bold"), relief=tk.FLAT, padx=15, pady=8, cursor="hand2")
        self.ai_apply_btn.pack(side=tk.LEFT, padx=4)
        self.ai_delete_btn = tk.Button(self.ai_action_frame, text="🗑️ 削除",
                                       command=self._delete_ai_proposal, bg="#e53e3e", fg="white",
                                       font=("Meiryo UI", 10, "bold"), relief=tk.FLAT, padx=15, pady=8, cursor="hand2")
        self.ai_delete_btn.pack(side=tk.LEFT, padx=4)
        self.ai_cancel_preview_btn = tk.Button(self.ai_action_frame, text="❌ 解除",
                                               command=self._cancel_preview, bg="#f56565", fg="white",
                                               font=("Meiryo UI", 10, "bold"), relief=tk.FLAT, padx=15, pady=8, cursor="hand2")
        self.ai_loading_label = tk.Label(main_container, text="", bg="#f5f7fa", fg="#718096",
                                         font=("Meiryo UI", 10))
        self._add_chat_message("system",
            "AIアシスタントが勤務表作成をお手伝いします。\n"
            "例：「休みを均等に配分して」「桜さんと小さんの当直を被らないように」など")
        if self.api_key:
            self._add_chat_message("system", "✅ APIキー設定済み — メッセージを送信できます。")
        else:
            if self._keyring_available:
                self._add_chat_message("system", "⚠️ APIキーが未設定です。右上の「🔑 APIキー」ボタンから設定してください。\n🔒 キーはOSの資格情報ストアに暗号化保存されます。")
            else:
                self._add_chat_message("system", "⚠️ APIキーが未設定です。右上の「🔑 APIキー」ボタンから設定してください。\n💡 pip install keyring で安全な保存が可能になります。")
        win.after(150, lambda: self.ai_input.focus_force())

    def _close_ai_window(self):
        if self.ai_window is not None:
            try:
                if self.ai_preview_active:
                    self._cancel_preview()
                self.ai_window.destroy()
            except tk.TclError:
                pass
            self.ai_window = None

    def _on_ai_enter_key(self, event):
        if event.state & 0x1:
            return
        self._send_ai_message()
        return "break"

    def _insert_example(self, text):
        if not hasattr(self, 'ai_input') or not self.ai_input.winfo_exists():
            return
        self.ai_input.delete("1.0", tk.END)
        self.ai_input.insert("1.0", text)
        self.ai_input.focus_force()

    def _reset_ai_conversation(self):
        if messagebox.askyesno("確認", "会話履歴をリセットしますか？"):
            self.ai_conversation_history.clear()
            self.ai_latest_proposal = None
            if hasattr(self, 'ai_chat_display') and self.ai_chat_display.winfo_exists():
                self.ai_chat_display.config(state=tk.NORMAL)
                self.ai_chat_display.delete("1.0", tk.END)
                self.ai_chat_display.config(state=tk.DISABLED)
                self._add_chat_message("system", "会話履歴がリセットされました。")
            self._hide_action_buttons()
            if self.ai_preview_active:
                self._cancel_preview()

    def _add_chat_message(self, role, content):
        if not hasattr(self, 'ai_chat_display') or not self.ai_chat_display.winfo_exists():
            return
        self.ai_chat_display.config(state=tk.NORMAL)
        if role == "user":
            self.ai_chat_display.insert(tk.END, "あなた:\n", "user")
            self.ai_chat_display.insert(tk.END, content + "\n\n", "")
        elif role == "assistant":
            self.ai_chat_display.insert(tk.END, "Claude:\n", "assistant")
            self.ai_chat_display.insert(tk.END, content + "\n\n", "")
        elif role == "system":
            self.ai_chat_display.insert(tk.END, "💡 " + content + "\n\n", "system")
        elif role == "proposal":
            self.ai_chat_display.insert(tk.END, "📋 提案内容:\n", "assistant")
            self.ai_chat_display.insert(tk.END, content + "\n\n", "proposal")
        self.ai_chat_display.see(tk.END)
        self.ai_chat_display.config(state=tk.DISABLED)

    def _send_ai_message(self):
        if not hasattr(self, 'ai_input') or not self.ai_input.winfo_exists():
            return
        if not self.api_key:
            self._add_chat_message("system", "⚠️ APIキーが設定されていません。「🔑 APIキー」ボタンから設定してください。")
            ok = self._prompt_api_key()
            if not ok or not self.api_key:
                return
            self._add_chat_message("system", "✅ APIキーが設定されました。")
        user_message = self.ai_input.get("1.0", tk.END).strip()
        if not user_message:
            return
        self.ai_input.delete("1.0", tk.END)
        self._add_chat_message("user", user_message)
        try:
            self.ai_send_btn.config(state=tk.DISABLED)
            self.ai_loading_label.config(text="🔄 Claudeが考えています...")
            self.ai_loading_label.pack(pady=10)
        except tk.TclError:
            pass
        threading.Thread(target=self._call_claude_api, args=(user_message,), daemon=True).start()

    def _call_claude_api(self, user_message):
        try:
            calendar_state = self._get_calendar_state()
            members_info = self._format_members_for_ai()
            team_rules_text = calendar_state['team_rules'] if calendar_state['team_rules'] else "特になし"
            system_prompt = f"""あなたは医療機関の勤務表作成支援AIです。
【全体の決まり事】
{team_rules_text}
【アクティブメンバー情報】
{members_info}
現在の状況:
- 年月: {self.current_year}年{self.current_month}月
- 月の日数: {calendar.monthrange(self.current_year, self.current_month)[1]}日
- 登録されている人員: {', '.join(calendar_state['people'])}
- 勤務種別: ER（4列）, A（4列）, B（4列）, 当直（4列）, 明け（4列）, 外勤（4列）, 出張（4列）, 休み（4列）
現在の勤務表データ:
{json.dumps(calendar_state['current_assignments'], ensure_ascii=False, indent=2)}
ロックされているセル（変更不可）:
{json.dumps(calendar_state['locked_cells'], ensure_ascii=False)}
ユーザーの要望に基づいて勤務表の提案を行ってください。
全体の決まり事とメンバーの個人リクエストを考慮してください。
**重要**: 提案する場合は、必ず以下のJSON形式で出力してください:
{{
  "explanation": "簡潔な説明（1-2文）",
  "assignments": [
    {{"day": 1, "label": "ER", "col": 0, "person": "桜"}},
    {{"day": 2, "label": "A", "col": 1, "person": "小"}}
  ]
}}"""
            self.ai_conversation_history.append({"role": "user", "content": user_message})
            import urllib.request
            import urllib.error
            data = {
                "model": "claude-sonnet-4-5",
                "max_tokens": 16000,
                "system": system_prompt,
                "messages": self.ai_conversation_history
            }
            req = urllib.request.Request(
                "https://api.anthropic.com/v1/messages",
                data=json.dumps(data).encode('utf-8'),
                headers={"Content-Type": "application/json", "x-api-key": self.api_key, "anthropic-version": "2023-06-01"}
            )
            with urllib.request.urlopen(req) as response:
                result = json.loads(response.read().decode('utf-8'))
            assistant_message = ""
            for content_block in result.get("content", []):
                if content_block.get("type") == "text":
                    assistant_message += content_block.get("text", "")
            self.ai_conversation_history.append({"role": "assistant", "content": assistant_message})
            self.root.after(0, self._handle_ai_response, assistant_message)
        except Exception as e:
            import urllib.error
            if hasattr(e, 'code'):
                if e.code == 401:
                    error_msg = "❌ APIキーが無効です。「🔑 APIキー」ボタンから正しいキーを設定してください。"
                    self.api_key = ""
                    self._delete_api_key()
                elif e.code == 429:
                    error_msg = "⏳ レート制限に達しました。少し待ってから再度お試しください。"
                else:
                    error_msg = f"❌ APIエラー（{e.code}）: {str(e)}"
            else:
                error_msg = f"❌ 通信エラー: {str(e)}"
            if self.ai_conversation_history and self.ai_conversation_history[-1]["role"] == "user":
                self.ai_conversation_history.pop()
            self.root.after(0, self._handle_ai_error, error_msg)

    def _handle_ai_response(self, response_text):
        if hasattr(self, 'ai_loading_label') and self.ai_loading_label.winfo_exists():
            self.ai_loading_label.pack_forget()
        if hasattr(self, 'ai_send_btn') and self.ai_send_btn.winfo_exists():
            self.ai_send_btn.config(state=tk.NORMAL)
        proposal = self._extract_json_from_response(response_text)
        if proposal and proposal.get("assignments"):
            explanation = proposal.get("explanation", "")
            self._add_chat_message("assistant", explanation)
            self._add_chat_message("proposal", f"📋 {len(proposal['assignments'])}件の割り当てを提案します")
            self.ai_latest_proposal = proposal
            self._show_action_buttons()
        else:
            self._add_chat_message("assistant", response_text)
            self._hide_action_buttons()

    def _handle_ai_error(self, error_msg):
        if hasattr(self, 'ai_loading_label') and self.ai_loading_label.winfo_exists():
            self.ai_loading_label.pack_forget()
        if hasattr(self, 'ai_send_btn') and self.ai_send_btn.winfo_exists():
            self.ai_send_btn.config(state=tk.NORMAL)
        self._add_chat_message("system", "❌ " + error_msg)

    def _extract_json_from_response(self, text):
        try:
            if "```json" in text:
                start = text.find("```json") + 7
                end = text.find("```", start)
                json_str = text[start:].strip() if end == -1 else text[start:end].strip()
            elif "```" in text:
                start = text.find("```") + 3
                end = text.find("```", start)
                json_str = text[start:].strip() if end == -1 else text[start:end].strip()
            else:
                start = text.find("{")
                if start == -1:
                    return None
                brace_count = 0
                end = start
                for i in range(start, len(text)):
                    if text[i] == "{": brace_count += 1
                    elif text[i] == "}":
                        brace_count -= 1
                        if brace_count == 0:
                            end = i + 1; break
                json_str = text[start:].strip() if end <= start else text[start:end]
            json_str = self._complete_incomplete_json(json_str)
            parsed = json.loads(json_str)
            if isinstance(parsed, dict) and "assignments" in parsed and isinstance(parsed.get("assignments"), list):
                return parsed
            return None
        except Exception:
            return None

    def _complete_incomplete_json(self, json_str):
        json_str = json_str.rstrip()
        if json_str.endswith(','):
            json_str = json_str[:-1].rstrip()
        open_braces = json_str.count('{')
        close_braces = json_str.count('}')
        open_brackets = json_str.count('[')
        close_brackets = json_str.count(']')
        missing_brackets = open_brackets - close_brackets
        missing_braces = open_braces - close_braces
        if missing_brackets > 0:
            json_str += '\n' + ']' * missing_brackets
        if missing_braces > 0:
            json_str += '\n' + '}' * missing_braces
        return json_str

    def _get_calendar_state(self):
        people = set()
        for value in self.cell_data.values():
            if value and value.strip():
                people.add(value)
        assignments = []
        for (day, label_idx, col), person in self.cell_data.items():
            if person and person.strip() and 1 <= label_idx <= 7:
                label = self.LABELS[label_idx]
                assignments.append({"day": day, "label": label, "col": col, "person": person})
        locked = []
        for (day, label_idx, col) in self.locked_cells:
            if 1 <= label_idx <= 7:
                label = self.LABELS[label_idx]
                person = self.cell_data.get((day, label_idx, col), "")
                locked.append({"day": day, "label": label, "col": col, "person": person})
        return {
            "people": sorted(people),
            "current_assignments": assignments,
            "locked_cells": locked,
            "team_rules": self.team_rules,
        }

    def _show_action_buttons(self):
        if hasattr(self, 'ai_action_frame') and self.ai_action_frame.winfo_exists():
            self.ai_action_frame.pack(fill=tk.X, pady=8)

    def _hide_action_buttons(self):
        if hasattr(self, 'ai_action_frame') and self.ai_action_frame.winfo_exists():
            self.ai_action_frame.pack_forget()

    def _preview_ai_proposal(self):
        if not self.ai_latest_proposal:
            return
        self.ai_preview_data = {'cell_data': copy.deepcopy(self.cell_data), 'cell_colors': copy.deepcopy(self.cell_colors)}
        proposed_people = set(a["person"] for a in self.ai_latest_proposal.get("assignments", []))
        for cell_key in [k for k, v in self.cell_data.items() if v in proposed_people]:
            if cell_key not in self.locked_cells:
                self.cell_data.pop(cell_key, None)
                self.cell_colors.pop(cell_key, None)
        for assignment in self.ai_latest_proposal.get("assignments", []):
            day = assignment["day"]; label = assignment["label"]; col = assignment["col"]; person = assignment["person"]
            try: label_idx = self.LABELS.index(label)
            except ValueError: continue
            cell_key = (day, label_idx, col)
            if cell_key in self.locked_cells: continue
            self.cell_data[cell_key] = person
            self.cell_colors[cell_key] = self._get_or_assign_color(person)
        self.ai_preview_active = True
        self._redraw_all_cells()
        try:
            self.ai_preview_btn.pack_forget()
            self.ai_apply_btn.pack_forget()
            self.ai_cancel_preview_btn.pack(side=tk.LEFT, padx=5)
        except tk.TclError:
            pass
        people_list = "、".join(sorted(proposed_people))
        self._add_chat_message("system", f"💡 プレビュー中です。\n対象メンバー: {people_list}\n「適用する」で確定、「プレビュー解除」で元に戻します。")

    def _cancel_preview(self):
        if not self.ai_preview_active:
            return
        self.cell_data = self.ai_preview_data['cell_data']
        self.cell_colors = self.ai_preview_data['cell_colors']
        self.ai_preview_data = {}
        self.ai_preview_active = False
        self._redraw_all_cells()
        try:
            self.ai_cancel_preview_btn.pack_forget()
            self.ai_preview_btn.pack(side=tk.LEFT, padx=5)
            self.ai_apply_btn.pack(side=tk.LEFT, padx=5)
        except tk.TclError:
            pass
        self._add_chat_message("system", "プレビューを解除しました。")

    def _apply_ai_proposal(self):
        if not self.ai_latest_proposal:
            return
        if self.ai_preview_active:
            self.ai_preview_data = {}
            self.ai_preview_active = False
            try:
                self.ai_cancel_preview_btn.pack_forget()
                self.ai_preview_btn.pack(side=tk.LEFT, padx=5)
                self.ai_apply_btn.pack(side=tk.LEFT, padx=5)
            except tk.TclError:
                pass
        else:
            self._save_state()
            proposed_people = set(a["person"] for a in self.ai_latest_proposal.get("assignments", []))
            for cell_key in [k for k, v in self.cell_data.items() if v in proposed_people]:
                if cell_key not in self.locked_cells:
                    self.cell_data.pop(cell_key, None)
                    self.cell_colors.pop(cell_key, None)
                    self.cell_sequence.pop(cell_key, None)
            for assignment in self.ai_latest_proposal.get("assignments", []):
                day = assignment["day"]; label = assignment["label"]; col = assignment["col"]; person = assignment["person"]
                try: label_idx = self.LABELS.index(label)
                except ValueError: continue
                cell_key = (day, label_idx, col)
                if cell_key in self.locked_cells: continue
                self.cell_data[cell_key] = person
                self.cell_colors[cell_key] = self._get_or_assign_color(person)
                if cell_key not in self.cell_sequence:
                    self.cell_sequence[cell_key] = self.next_sequence
                    self.next_sequence += 1
            self._redraw_all_cells()
        self._add_chat_message("system", "✅ 提案を適用しました！")
        self._hide_action_buttons()
        self.ai_latest_proposal = None
        self.notebook.select(0)
        self.draw_statistics()

    def _delete_ai_proposal(self):
        if not self.ai_latest_proposal:
            return
        if self.ai_preview_active:
            self._cancel_preview()
        self.ai_latest_proposal = None
        self._hide_action_buttons()
        self._add_chat_message("system", "🗑️ 提案を削除しました。")

    def _on_tab_changed(self, event):
        current_tab = self.notebook.index(self.notebook.select())
        if current_tab == 1:
            self.draw_sub_calendar()
        elif current_tab == 2:
            self.draw_personal_calendars()

    def _open_special_item_window(self, category: str):
        win = self._special_item_windows.get(category)
        if win and tk.Toplevel.winfo_exists(win):
            win.lift(); win.focus_force(); return
        icons = {"外勤": "🚗", "委員会": "🏛", "コース": "📚", "訓練": "🎯"}
        icon = icons.get(category, "📋")
        win = tk.Toplevel(self.root)
        win.title(f"{icon} {category} 項目管理")
        win.geometry("420x480")
        win.resizable(True, True)
        win.configure(bg="#f5f7fa")
        win.transient(self.root)
        self._special_item_windows[category] = win
        header = tk.Frame(win, bg="#2d5a7b", pady=10)
        header.pack(fill=tk.X)
        tk.Label(header, text=f"{icon} {category} 項目一覧", font=("Meiryo UI", 13, "bold"),
                 bg="#2d5a7b", fg="white").pack()
        tk.Label(header, text="右クリックメニューに表示されます",
                 font=("Meiryo UI", 9), bg="#2d5a7b", fg="#aac8e0").pack()
        list_frame = tk.Frame(win, bg="#f5f7fa")
        list_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=(12, 5))
        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        listbox = tk.Listbox(list_frame, font=("Meiryo UI", 11), yscrollcommand=scrollbar.set,
                             selectmode=tk.SINGLE, bg="white", bd=1, relief=tk.SOLID,
                             activestyle="none", selectbackground="#2d5a7b", selectforeground="white")
        listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=listbox.yview)
        def refresh_list():
            listbox.delete(0, tk.END)
            for item in self.special_items[category]:
                listbox.insert(tk.END, f"  {item}")
        refresh_list()
        input_frame = tk.Frame(win, bg="#f5f7fa")
        input_frame.pack(fill=tk.X, padx=15, pady=5)
        tk.Label(input_frame, text="新しい項目名:", font=("Meiryo UI", 10), bg="#f5f7fa").pack(anchor=tk.W)
        entry_frame = tk.Frame(input_frame, bg="#f5f7fa")
        entry_frame.pack(fill=tk.X, pady=3)
        entry = tk.Entry(entry_frame, font=("Meiryo UI", 11), bd=1, relief=tk.SOLID)
        entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 6))
        def add_item():
            text = entry.get().strip()
            if not text: return
            if text in self.special_items[category]:
                messagebox.showwarning("重複", f"「{text}」は既に登録されています", parent=win); return
            self.special_items[category].append(text)
            refresh_list(); entry.delete(0, tk.END); listbox.see(tk.END)
        tk.Button(entry_frame, text="追加", command=add_item, bg="#2d5a7b", fg="white",
                  font=("Meiryo UI", 10, "bold"), width=7, relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT)
        entry.bind("<Return>", lambda e: add_item())
        btn_frame = tk.Frame(win, bg="#f5f7fa")
        btn_frame.pack(fill=tk.X, padx=15, pady=(5, 15))
        def delete_item():
            sel = listbox.curselection()
            if not sel: messagebox.showinfo("未選択", "削除する項目を選択してください", parent=win); return
            idx = sel[0]
            if messagebox.askyesno("確認", f"「{self.special_items[category][idx]}」を削除しますか？", parent=win):
                self.special_items[category].pop(idx); refresh_list()
        def move_up():
            sel = listbox.curselection()
            if not sel or sel[0] == 0: return
            idx = sel[0]; items = self.special_items[category]
            items[idx-1], items[idx] = items[idx], items[idx-1]
            refresh_list(); listbox.selection_set(idx-1)
        def move_down():
            sel = listbox.curselection()
            if not sel: return
            idx = sel[0]; items = self.special_items[category]
            if idx >= len(items) - 1: return
            items[idx], items[idx+1] = items[idx+1], items[idx]
            refresh_list(); listbox.selection_set(idx+1)
        tk.Button(btn_frame, text="🗑 削除", command=delete_item, bg="#e76f51", fg="white",
                  font=("Meiryo UI", 10), width=9, relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT, padx=(0, 4))
        tk.Button(btn_frame, text="▲ 上へ", command=move_up, bg="#6c757d", fg="white",
                  font=("Meiryo UI", 10), width=8, relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT, padx=2)
        tk.Button(btn_frame, text="▼ 下へ", command=move_down, bg="#6c757d", fg="white",
                  font=("Meiryo UI", 10), width=8, relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT, padx=2)
        tk.Button(btn_frame, text="✓ 閉じる", command=win.destroy, bg="#2a9d8f", fg="white",
                  font=("Meiryo UI", 10), width=9, relief=tk.FLAT, cursor="hand2").pack(side=tk.RIGHT)
        win.focus_force()

    def switch_stats_mode(self, mode):
        self.stats_mode = mode
        if mode == "person":
            self.btn_person.config(relief=tk.SUNKEN, bg="#4a90e2", font=("Meiryo UI", 10, "bold"))
            self.btn_category.config(relief=tk.RAISED, bg="#7c8a9e", font=("Meiryo UI", 10))
        else:
            self.btn_category.config(relief=tk.SUNKEN, bg="#4a90e2", font=("Meiryo UI", 10, "bold"))
            self.btn_person.config(relief=tk.RAISED, bg="#7c8a9e", font=("Meiryo UI", 10))
        self.draw_statistics()

    def _toggle_sound(self):
        is_muted = self.sound_manager.toggle_mute()
        if is_muted:
            self.sound_btn.config(text="🔇 SE:OFF", bg="#9E9E9E")
        else:
            self.sound_btn.config(text="🔊 SE:ON", bg="#4CAF50")
            self.sound_manager.play("bright")

    def _on_close(self):
        self._close_ai_window()
        self.sound_manager.cleanup()
        self.root.destroy()

    def _mark_day_as_complete(self, day: int):
        if day in self.completed_days: return
        self.completed_days.add(day)
        self._draw_day_complete_marker(day)
        self.sound_manager.play("complete")
        self._animate_completion(day)

    def _unmark_day_as_complete(self, day: int):
        if day not in self.completed_days: return
        self.completed_days.discard(day)
        if day in self.day_complete_rects:
            self.canvas.delete(self.day_complete_rects[day])
            del self.day_complete_rects[day]

    def _draw_day_complete_marker(self, day: int):
        if day not in self._day_positions: return
        week, dow, col_x, row_y = self._day_positions[day]
        x1 = col_x; y1 = row_y
        x2 = col_x + self.BLOCK_W * self.CELL_W; y2 = row_y + self.BLOCK_H * self.CELL_H
        rect_id = self.canvas.create_rectangle(x1, y1, x2, y2, fill="#fffacd", outline="", tags="day_complete")
        self.day_complete_rects[day] = rect_id
        self.canvas.tag_lower("day_complete")
        self.canvas.tag_lower("day_complete", "grid")

    def _redraw_day_complete_markers(self):
        for rect_id in self.day_complete_rects.values():
            self.canvas.delete(rect_id)
        self.day_complete_rects.clear()
        for day in list(self.completed_days):
            self._draw_day_complete_marker(day)

    def _animate_completion(self, day: int):
        info = self._day_positions.get(day)
        if not info: return
        week, dow, col_x, row_y = info
        x1 = col_x; y1 = row_y
        x2 = col_x + self.BLOCK_W * self.CELL_W
        y2 = row_y + self.BLOCK_H * self.CELL_H
        animation_ids = []
        def phase1():
            border_id = self.canvas.create_rectangle(x1, y1, x2, y2, outline="#FFD700", width=6, tags="completion_anim")
            animation_ids.append(border_id)
            self.root.after(100, phase2)
        def phase2():
            for _ in range(12):
                star_x = random.randint(int(x1+10), int(x2-10))
                star_y = random.randint(int(y1+10), int(y2-10))
                star_id = self.canvas.create_text(star_x, star_y, text="✨",
                    font=("Segoe UI Emoji", random.randint(12, 20)),
                    fill=random.choice(["#FFD700","#FFA500","#FFFF00","#FF69B4"]), tags="completion_anim")
                animation_ids.append(star_id)
            self.root.after(150, phase3)
        def phase3():
            for item_id in animation_ids:
                if self.canvas.type(item_id) == "rectangle":
                    self.canvas.itemconfig(item_id, width=4, outline="#FFD700")
            self.root.after(150, phase4)
        def phase4():
            for _ in range(8):
                star_x = random.randint(int(x1+10), int(x2-10))
                star_y = random.randint(int(y1+10), int(y2-10))
                star_id = self.canvas.create_text(star_x, star_y, text="⭐",
                    font=("Segoe UI Emoji", random.randint(10, 18)),
                    fill=random.choice(["#FFD700","#FFA500","#FFFF00"]), tags="completion_anim")
                animation_ids.append(star_id)
            self.root.after(200, phase5)
        def phase5():
            for item_id in animation_ids:
                if self.canvas.type(item_id) == "rectangle":
                    self.canvas.itemconfig(item_id, width=7, outline="#FFA500")
            self.root.after(120, phase6)
        def phase6():
            for item_id in animation_ids:
                if self.canvas.type(item_id) == "rectangle":
                    self.canvas.itemconfig(item_id, width=5, outline="#FFD700")
            self.root.after(150, phase7)
        def phase7():
            for item_id in animation_ids:
                if self.canvas.type(item_id) == "rectangle":
                    self.canvas.itemconfig(item_id, width=3, outline="#FFE4B5")
            self.root.after(180, phase8)
        def phase8():
            self.canvas.delete("completion_anim")
        phase1()

    def _open_day_memo(self, day: int):
        if day in self.memo_windows:
            try:
                self.memo_windows[day].lift()
                self.memo_windows[day].focus_force()
                return
            except tk.TclError:
                del self.memo_windows[day]
        if len(self.memo_windows) >= 10:
            messagebox.showwarning("制限", "メモ帳は最大10枚までしか同時に開けません。\n他のメモ帳を閉じてから開いてください。")
            return
        memo_win = tk.Toplevel(self.root)
        memo_win.title(f"📝 {day}日")
        memo_width = 220; memo_height = 140
        memo_win.configure(bg="#f8f9fa")
        header = tk.Frame(memo_win, bg="#4a90e2", height=28)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        tk.Label(header, text=f"📅 {self.current_month}/{day}", bg="#4a90e2", fg="white",
                 font=("Meiryo UI", 9, "bold")).pack(side=tk.LEFT, padx=8, pady=4)
        text_frame = tk.Frame(memo_win, bg="#f8f9fa")
        text_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        text_widget = tk.Text(text_frame, height=2, font=("Meiryo UI", 9), bg="white", fg="#2d3748",
                              relief=tk.SOLID, bd=1, wrap=tk.WORD)
        text_widget.pack(fill=tk.BOTH, expand=True)
        if day in self.day_memos:
            text_widget.insert("1.0", self.day_memos[day])
        close_btn = tk.Button(header, text="✕",
                              command=lambda: self._close_day_memo(day, memo_win, text_widget),
                              bg="#e74c3c", fg="white", font=("Meiryo UI", 8, "bold"),
                              relief=tk.FLAT, cursor="hand2", padx=4, pady=0)
        close_btn.pack(side=tk.RIGHT, padx=5)
        memo_win.update_idletasks()
        self.root.update_idletasks()
        try:
            main_x = self.root.winfo_x(); main_y = self.root.winfo_y()
            main_width = self.root.winfo_width()
            memo_x = main_x + main_width + 10
            memo_y = main_y + 100 + (len(self.memo_windows) * 150)
            if day in self._day_positions:
                week, dow, col_x, row_y = self._day_positions[day]
                try:
                    canvas_x = self.canvas.winfo_rootx(); canvas_y = self.canvas.winfo_rooty()
                    if canvas_x > 0 and canvas_y > 0:
                        block_width = int(self.BLOCK_W * self.CELL_W)
                        memo_x = canvas_x + int(col_x) + block_width + 10
                        memo_y = canvas_y + int(row_y)
                except: pass
        except:
            screen_width = self.root.winfo_screenwidth(); screen_height = self.root.winfo_screenheight()
            memo_x = (screen_width - memo_width) // 2; memo_y = (screen_height - memo_height) // 2
        memo_win.geometry(f"{memo_width}x{memo_height}+{memo_x}+{memo_y}")
        memo_win.resizable(False, False)
        memo_win.lift()
        memo_win.attributes('-topmost', True)
        text_widget.focus_force()
        self.memo_windows[day] = memo_win
        memo_win.protocol("WM_DELETE_WINDOW", lambda: self._close_day_memo(day, memo_win, text_widget))

    def _close_day_memo(self, day: int, window: tk.Toplevel, text_widget: tk.Text):
        memo_text = text_widget.get("1.0", "end-1c").strip()
        if memo_text:
            self.day_memos[day] = memo_text
        else:
            if day in self.day_memos:
                del self.day_memos[day]
        if day in self.memo_windows:
            del self.memo_windows[day]
        try: window.destroy()
        except: pass


    def create_calendar(self):
        try:
            self.current_year = int(self.year_var.get())
            self.current_month = int(self.month_var.get())
        except ValueError:
            messagebox.showerror("エラー", "正しい年月を入力してください")
            return
        # ✨ 改善: ウィンドウタイトルに年月を表示
        self.root.title(f"カレンダー管理システム — {self.current_year}年{self.current_month}月")
        # ✨ 改善: 祝日を取得
        self._current_holidays = self._get_japanese_holidays(self.current_year, self.current_month)
        self._finish_edit()
        self.canvas.delete("all")
        self._rect_ids.clear(); self._text_ids.clear(); self._cell_bounds.clear()
        self._day_positions.clear(); self.day_complete_rects.clear()
        self.left_frame.update_idletasks()
        canvas_w = self.canvas.winfo_width(); canvas_h = self.canvas.winfo_height()
        if canvas_w < 100: canvas_w = 1000
        if canvas_h < 100: canvas_h = 800
        first_wd = calendar.weekday(self.current_year, self.current_month, 1)
        first_wd = (first_wd + 1) % 7
        if first_wd == 0: first_wd = 7
        last_day = calendar.monthrange(self.current_year, self.current_month)[1]
        if self.current_month == 1:
            prev_month_year = self.current_year - 1; prev_month = 12
        else:
            prev_month_year = self.current_year; prev_month = self.current_month - 1
        prev_month_last_day = calendar.monthrange(prev_month_year, prev_month)[1]
        total_cells = first_wd - 1 + last_day
        total_weeks = -(-total_cells // 7)
        ox = 10; oy = 10
        scrollbar_buffer = 20
        available_w = canvas_w - (ox * 2) - (self.LABEL_COL_W * 2) - scrollbar_buffer
        total_cols = 7 * self.BLOCK_W
        new_cell_w = available_w / total_cols
        available_h = canvas_h - oy - self.TITLE_H - self.HEADER_H - scrollbar_buffer
        total_rows = total_weeks * self.BLOCK_H
        new_cell_h = available_h / total_rows
        self.CELL_W = max(20, new_cell_w); self.CELL_H = max(15, new_cell_h)
        self.canvas.create_text(ox, oy, anchor="nw",
                                text=f"{self.current_year}年 {self.current_month}月",
                                font=("Meiryo UI", 24, "bold"), fill="black")
        oy += self.TITLE_H
        hdr_x = ox + self.LABEL_COL_W
        for i, wd in enumerate(self.WEEKDAYS):
            x1 = hdr_x + i * self.BLOCK_W * self.CELL_W
            x2 = x1 + self.BLOCK_W * self.CELL_W
            self.canvas.create_rectangle(x1, oy, x2, oy + self.HEADER_H, fill="#dcdcdc", outline="gray")
            fg = "blue" if i == 5 else ("red" if i == 6 else "black")
            self.canvas.create_text((x1 + x2) / 2, oy + self.HEADER_H / 2,
                                    text=wd, font=("Meiryo UI", 12, "bold"), fill=fg)
        oy += self.HEADER_H
        current_day = 1; prev_day_counter = first_wd - 2; next_day_counter = 1
        right_x = ox + self.LABEL_COL_W + 7 * self.BLOCK_W * self.CELL_W + self.LABEL_COL_W
        for week in range(total_weeks):
            row_y = oy + week * self.BLOCK_H * self.CELL_H
            self._draw_label_column(ox, row_y)
            for dow in range(7):
                col_x = ox + self.LABEL_COL_W + dow * self.BLOCK_W * self.CELL_W
                if week == 0 and dow < first_wd - 1:
                    prev_day = prev_month_last_day - prev_day_counter
                    self._draw_adjacent_month_block(col_x, row_y, -prev_day, dow, week, is_prev_month=True)
                    prev_day_counter -= 1
                elif current_day <= last_day:
                    self._draw_day_block(col_x, row_y, current_day, dow, week)
                    current_day += 1
                else:
                    self._draw_adjacent_month_block(col_x, row_y, last_day + next_day_counter, dow, week, is_prev_month=False)
                    next_day_counter += 1
                if dow < 6:
                    line_x = col_x + self.BLOCK_W * self.CELL_W
                    self.canvas.create_line(line_x, row_y, line_x, row_y + self.BLOCK_H * self.CELL_H,
                                          width=3, fill="#444444", tags="week_separator")
            right_label_x = ox + self.LABEL_COL_W + 7 * self.BLOCK_W * self.CELL_W
            self._draw_label_column(right_label_x, row_y)
            week_bottom_y = row_y + self.BLOCK_H * self.CELL_H
            self.canvas.create_line(ox, week_bottom_y, right_x, week_bottom_y, width=3, fill="#444444", tags="grid_line")
            self.canvas.tag_raise("divider_line")
        total_h = oy + total_weeks * self.BLOCK_H * self.CELL_H + 20
        total_w = ox + self.LABEL_COL_W * 2 + 7 * self.BLOCK_W * self.CELL_W + 20
        self.canvas.configure(scrollregion=(0, 0, total_w, total_h))
        self._apply_data()
        self._draw_all_lines()
        self.canvas.tag_raise("divider_line")
        self.canvas.update_idletasks()
        self.canvas.tag_raise("divider_line")
        self._redraw_special_borders()
        self._redraw_lock_borders()
        self._redraw_day_complete_markers()
        self.draw_statistics()


    def draw_statistics(self):
        self.right_frame.update_idletasks()
        self.stats_canvas.delete("all")
        w = self.stats_canvas.winfo_width()
        if w < 50: w = 350
        valid_values = [v for v in self.cell_data.values() if v and v.strip()]
        if not valid_values:
            self.stats_canvas.create_text(w/2, 50, text="データなし", font=("Meiryo UI", 11))
            self.stats_canvas.configure(scrollregion=(0, 0, w, 100))
            return
        if self.stats_mode == "person":
            self._draw_person_statistics(w)
        else:
            self._draw_category_statistics(w)

    def _draw_person_statistics(self, canvas_width):
        w = canvas_width
        person_stats = defaultdict(lambda: defaultdict(int))
        for (day, label_idx, col_idx), value in self.cell_data.items():
            if not value or not value.strip(): continue
            label = self.LABELS[label_idx] if label_idx < len(self.LABELS) else ""
            if label and label.strip():
                person_stats[value][label] += 1
        if not person_stats:
            self.stats_canvas.create_text(w/2, 50, text="データなし", font=("Meiryo UI", 11))
            self.stats_canvas.configure(scrollregion=(0, 0, w, 100))
            return
        sorted_persons = sorted(person_stats.keys(), key=lambda p: sum(person_stats[p].values()), reverse=True)
        y_pos = 15
        self.stats_canvas.create_text(w/2, y_pos, anchor="n", text="【詳細ビュー - 個人別内訳】",
                                      font=("Meiryo UI", 13, "bold", "underline"), fill="#2c3e50")
        y_pos += 30
        category_order = ["ER","A","B","当直","明け","外勤","出張","休み"]
        self._hp_tooltip_areas = []
        for person in sorted_persons:
            stats = person_stats[person]
            work_days = sum(count for cat, count in stats.items() if cat != "休み")
            holiday_days = stats.get("休み", 0)
            total_days = work_days + holiday_days
            hp_pct, fatigue_detail = self._calc_fatigue(person)
            self.stats_canvas.create_rectangle(10, y_pos-2, w-10, y_pos+18, fill="#ecf0f1", outline="#bdc3c7")
            self.stats_canvas.create_text(15, y_pos+8, anchor="w",
                text=f"■ {person}  総計: {total_days}日 （労働: {work_days}日、休み: {holiday_days}日）",
                font=("Meiryo UI", 10, "bold"))
            y_pos += 20
            self.stats_canvas.create_rectangle(10, y_pos, w-10, y_pos+20, fill="#dde4ea", outline="#bdc3c7")
            gauge_w = w - 80; gauge_x = 40; gauge_y = y_pos + 10
            self._draw_hp_gauge(gauge_x, gauge_y, gauge_w, hp_pct)
            tooltip_text = self._build_hp_tooltip_text(person, hp_pct, fatigue_detail)
            self._hp_tooltip_areas.append((10, y_pos, w-10, y_pos+20, tooltip_text))
            y_pos += 26
            max_count = max([stats.get(cat, 0) for cat in category_order]) if stats else 1
            bar_max_width = w - 140
            for category in category_order:
                count = stats.get(category, 0)
                color = self._get_or_assign_color(category)
                self.stats_canvas.create_text(30, y_pos, anchor="w", text=category, font=("Meiryo UI", 9))
                if count > 0:
                    bar_width = (count / max(max_count, 1)) * bar_max_width
                    self.stats_canvas.create_rectangle(90, y_pos-8, 90+bar_width, y_pos+8, fill=color, outline=color, width=0)
                    self.stats_canvas.create_text(95+bar_width, y_pos, anchor="w", text=f"{count}日", font=("Meiryo UI", 9, "bold"))
                else:
                    self.stats_canvas.create_text(90, y_pos, anchor="w", text="─ 0日", font=("Meiryo UI", 9), fill="#95a5a6")
                y_pos += 20
            y_pos += 5
            self.stats_canvas.create_line(10, y_pos, w-10, y_pos, fill="#bdc3c7", width=1)
            y_pos += 15
        y_pos += 10
        self.stats_canvas.create_text(w/2, y_pos, anchor="n", text="【比較ビュー - 項目別人員比較】",
                                      font=("Meiryo UI", 13, "bold", "underline"), fill="#2c3e50")
        y_pos += 35
        person_col_width = 70; start_x = 80
        for i, person in enumerate(sorted_persons):
            x_pos = start_x + i * person_col_width
            self.stats_canvas.create_text(x_pos + person_col_width/2, y_pos, anchor="n", text=person, font=("Meiryo UI", 9, "bold"))
        y_pos += 25
        for category in category_order:
            color = self._get_or_assign_color(category)
            self.stats_canvas.create_text(15, y_pos, anchor="w", text=category, font=("Meiryo UI", 9, "bold"))
            all_counts = [person_stats[p].get(category, 0) for p in sorted_persons]
            max_count = max(all_counts) if all_counts else 1
            for i, person in enumerate(sorted_persons):
                count = person_stats[person].get(category, 0)
                x_pos = start_x + i * person_col_width
                if count > 0:
                    bar_width = (count / max(max_count, 1)) * (person_col_width - 25)
                    self.stats_canvas.create_rectangle(x_pos, y_pos-7, x_pos+bar_width, y_pos+7, fill=color, outline="", width=0)
                    self.stats_canvas.create_text(x_pos+bar_width+3, y_pos, anchor="w", text=str(count), font=("Meiryo UI", 8, "bold"))
                else:
                    self.stats_canvas.create_text(x_pos, y_pos, anchor="w", text="─", font=("Meiryo UI", 9), fill="#95a5a6")
            y_pos += 22
        total_width = max(w, start_x + len(sorted_persons) * person_col_width + 50)
        self.stats_canvas.configure(scrollregion=(0, 0, total_width, y_pos + 20))

    def _calc_fatigue(self, person: str):
        person_days: Dict[int, set] = {}
        for (day, li, ci), name in self.cell_data.items():
            if name != person: continue
            label = self.LABELS[li] if li < len(self.LABELS) else ""
            if label and label.strip():
                person_days.setdefault(day, set()).add(label)
        tochoku = sum(1 for ls in person_days.values() if "当直" in ls)
        kyukin  = sum(1 for ls in person_days.values() if "休み" in ls)
        gakin   = sum(1 for ls in person_days.values() if "外勤" in ls)
        shutcho = sum(1 for ls in person_days.values() if "出張" in ls)
        work_days = sorted(d for d, ls in person_days.items() if "休み" not in ls)
        max_consec = cur = 0; prev = None
        for d in work_days:
            cur = cur + 1 if prev is not None and d == prev + 1 else 1
            max_consec = max(max_consec, cur); prev = d
        score = tochoku*15 + max(0, max_consec-2)*8 + gakin*5 + shutcho*8 - kyukin*10
        score = max(0, score)
        MAX_SCORE = 240
        fatigue_pct = min(100, int(score / MAX_SCORE * 100))
        hp_pct = 100 - fatigue_pct
        detail = {"当直回数": tochoku, "休み回数": kyukin, "外勤回数": gakin, "出張回数": shutcho, "連続勤務最長": max_consec, "疲労スコア": score}
        return hp_pct, detail

    def _draw_hp_gauge(self, x: float, cy: float, total_w: float, hp_pct: int) -> list:
        ids = []
        if hp_pct >= 80: bar_color, emoji = "#27ae60", "😊"
        elif hp_pct >= 60: bar_color, emoji = "#f1c40f", "😐"
        elif hp_pct >= 40: bar_color, emoji = "#e67e22", "😓"
        elif hp_pct >= 20: bar_color, emoji = "#e74c3c", "😫"
        else: bar_color, emoji = "#900000", "💀"
        bar_h = 12; y1 = cy - bar_h//2; y2 = cy + bar_h//2
        ids.append(self.stats_canvas.create_text(x-2, cy, anchor="w", text="疲労度", font=("Meiryo UI", 8), fill="#555"))
        label_prefix_w = 34; bar_w = total_w - label_prefix_w - 90; bx = x + label_prefix_w
        ids.append(self.stats_canvas.create_rectangle(bx, y1, bx+bar_w, y2, fill="#c8ccd4", outline="#999", width=1))
        filled_w = max(1, int(bar_w * hp_pct / 100))
        ids.append(self.stats_canvas.create_rectangle(bx, y1, bx+filled_w, y2, fill=bar_color, outline="", width=0))
        ids.append(self.stats_canvas.create_text(bx+bar_w+4, cy, anchor="w", text=f"HP {hp_pct}%", font=("Meiryo UI", 8, "bold"), fill=bar_color))
        ids.append(self.stats_canvas.create_text(bx+bar_w+62, cy, anchor="center", text=emoji, font=("Meiryo UI", 13)))
        return ids

    def _build_hp_tooltip_text(self, person: str, hp_pct: int, detail: dict) -> str:
        if hp_pct >= 80: comment = "十分な余裕あり。良いバランスです！"
        elif hp_pct >= 60: comment = "やや負荷あり。休みを確保したいところ。"
        elif hp_pct >= 40: comment = "疲労が蓄積ぎみ。注意が必要です。"
        elif hp_pct >= 20: comment = "かなりの過負荷状態！要フォロー。"
        else: comment = "燃え尽き寸前…！至急休みを。"
        lines = [f"── {person} の疲労度詳細 ──", f"  HP: {hp_pct}%  {comment}", "",
                 f"  当直回数     : {detail['当直回数']} 回  (+{detail['当直回数']*15}pt)",
                 f"  外勤回数     : {detail['外勤回数']} 回  (+{detail['外勤回数']*5}pt)",
                 f"  出張回数     : {detail['出張回数']} 回  (+{detail['出張回数']*8}pt)",
                 f"  連続勤務最長 : {detail['連続勤務最長']} 日"]
        if detail['連続勤務最長'] > 2:
            lines[-1] += f"  (+{max(0,detail['連続勤務最長']-2)*8}pt)"
        lines += [f"  休み回数     : {detail['休み回数']} 回  (-{detail['休み回数']*10}pt)", "",
                  f"  疲労スコア合計: {detail['疲労スコア']} pt"]
        return "\n".join(lines)

    def _on_stats_motion(self, event):
        cx = self.stats_canvas.canvasx(event.x); cy = self.stats_canvas.canvasy(event.y)
        for (x1, y1, x2, y2, text) in self._hp_tooltip_areas:
            if x1 <= cx <= x2 and y1 <= cy <= y2:
                self._show_hp_tooltip(event, text); return
        self._hide_hp_tooltip()

    def _show_hp_tooltip(self, event, text: str):
        self._hide_hp_tooltip()
        win = tk.Toplevel(self.root)
        win.wm_overrideredirect(True); win.wm_attributes("-topmost", True)
        win.geometry(f"+{event.x_root+12}+{event.y_root+12}")
        frame = tk.Frame(win, background="#fffbe6", highlightbackground="#e0c050", highlightthickness=1)
        frame.pack()
        tk.Label(frame, text=text, background="#fffbe6", font=("Meiryo UI", 9), justify=tk.LEFT, padx=8, pady=6).pack()
        self._hp_tooltip_win = win

    def _hide_hp_tooltip(self, event=None):
        if self._hp_tooltip_win is not None:
            try: self._hp_tooltip_win.destroy()
            except Exception: pass
            self._hp_tooltip_win = None

    def _draw_category_statistics(self, canvas_width):
        w = canvas_width
        category_stats = defaultdict(lambda: defaultdict(int))
        for (day, label_idx, col_idx), value in self.cell_data.items():
            if not value or not value.strip(): continue
            label = self.LABELS[label_idx] if label_idx < len(self.LABELS) else ""
            if label and label.strip():
                category_stats[label][value] += 1
        if not category_stats:
            self.stats_canvas.create_text(w/2, 50, text="データなし", font=("Meiryo UI", 11))
            self.stats_canvas.configure(scrollregion=(0, 0, w, 100)); return
        priority_order = ["ER","A","B","当直","明け","外勤","出張","休み"]
        sorted_categories = [cat for cat in priority_order if cat in category_stats]
        sorted_categories.extend([cat for cat in category_stats.keys() if cat not in priority_order])
        all_persons = set()
        for persons_dict in category_stats.values():
            all_persons.update(persons_dict.keys())
        sorted_persons = sorted(all_persons)
        y_pos = 15
        self.stats_canvas.create_text(w/2, y_pos, anchor="n", text="【詳細ビュー - 項目別内訳】",
                                      font=("Meiryo UI", 13, "bold", "underline"), fill="#2c3e50")
        y_pos += 30
        for category in sorted_categories:
            persons = category_stats[category]; total = sum(persons.values())
            color = self._get_or_assign_color(category)
            self.stats_canvas.create_rectangle(10, y_pos-2, w-10, y_pos+20, fill="#ecf0f1", outline="#bdc3c7")
            self.stats_canvas.create_text(15, y_pos+9, anchor="w", text=f"■ {category}  総計: {total}日", font=("Meiryo UI", 10, "bold"))
            y_pos += 26
            max_count = max(persons.values()) if persons else 1; bar_max_width = w - 140
            for person in sorted_persons:
                count = persons.get(person, 0); person_color = self._get_or_assign_color(person)
                self.stats_canvas.create_text(30, y_pos, anchor="w", text=person, font=("Meiryo UI", 9))
                if count > 0:
                    bar_width = (count / max(max_count, 1)) * bar_max_width
                    self.stats_canvas.create_rectangle(90, y_pos-8, 90+bar_width, y_pos+8, fill=person_color, outline="", width=0)
                    percentage = (count / total * 100) if total > 0 else 0
                    self.stats_canvas.create_text(95+bar_width, y_pos, anchor="w", text=f"{count}日 ({percentage:.1f}%)", font=("Meiryo UI", 9, "bold"))
                else:
                    self.stats_canvas.create_text(90, y_pos, anchor="w", text="─ 0日", font=("Meiryo UI", 9), fill="#95a5a6")
                y_pos += 20
            y_pos += 5
            self.stats_canvas.create_line(10, y_pos, w-10, y_pos, fill="#bdc3c7", width=1)
            y_pos += 15
        y_pos += 10
        self.stats_canvas.create_text(w/2, y_pos, anchor="n", text="【比較ビュー - 人別項目比較】",
                                      font=("Meiryo UI", 13, "bold", "underline"), fill="#2c3e50")
        y_pos += 35
        cat_col_width = 70; start_x = 80
        for i, category in enumerate(sorted_categories):
            x_pos = start_x + i * cat_col_width
            self.stats_canvas.create_text(x_pos+cat_col_width/2, y_pos, anchor="n", text=category, font=("Meiryo UI", 9, "bold"))
        y_pos += 25
        for person in sorted_persons:
            person_color = self._get_or_assign_color(person)
            self.stats_canvas.create_text(15, y_pos, anchor="w", text=person, font=("Meiryo UI", 9, "bold"))
            all_counts = [category_stats[cat].get(person, 0) for cat in sorted_categories]
            max_count = max(all_counts) if all_counts else 1
            for i, category in enumerate(sorted_categories):
                count = category_stats[category].get(person, 0)
                x_pos = start_x + i * cat_col_width
                color = self._get_or_assign_color(category)
                if count > 0:
                    bar_width = (count / max(max_count, 1)) * (cat_col_width - 25)
                    self.stats_canvas.create_rectangle(x_pos, y_pos-7, x_pos+bar_width, y_pos+7, fill=color, outline="", width=0)
                    self.stats_canvas.create_text(x_pos+bar_width+3, y_pos, anchor="w", text=str(count), font=("Meiryo UI", 8, "bold"))
                else:
                    self.stats_canvas.create_text(x_pos, y_pos, anchor="w", text="─", font=("Meiryo UI", 9), fill="#95a5a6")
            y_pos += 22
        total_width = max(w, start_x + len(sorted_categories) * cat_col_width + 50)
        self.stats_canvas.configure(scrollregion=(0, 0, total_width, y_pos + 20))


    def draw_personal_calendars(self):
        self.personal_canvas.delete("all")
        all_persons = set()
        for value in self.cell_data.values():
            if value and value.strip(): all_persons.add(value)
        if not all_persons:
            self.personal_canvas.create_text(400, 300, text="データがありません", font=("Meiryo UI", 16))
            self.personal_canvas.configure(scrollregion=(0, 0, 800, 600)); return
        sorted_persons = sorted(all_persons)
        cols = 3; cal_width = 300; cal_height = 280; margin = 30
        start_x = margin; start_y = margin
        first_wd = calendar.weekday(self.current_year, self.current_month, 1)
        first_wd = (first_wd + 1) % 7
        if first_wd == 0: first_wd = 7
        last_day = calendar.monthrange(self.current_year, self.current_month)[1]
        for person_idx, person in enumerate(sorted_persons):
            row = person_idx // cols; col = person_idx % cols
            base_x = start_x + col * (cal_width + margin)
            base_y = start_y + row * (cal_height + margin)
            self.personal_canvas.create_text(base_x + cal_width/2, base_y,
                text=f"{person} の勤務カレンダー", font=("Meiryo UI", 12, "bold"))
            cell_size = 35; grid_start_x = base_x + 10; grid_start_y = base_y + 30
            weekdays_short = ["月","火","水","木","金","土","日"]
            for wd_idx, wd in enumerate(weekdays_short):
                x = grid_start_x + wd_idx * cell_size; y = grid_start_y
                self.personal_canvas.create_text(x+cell_size/2, y+cell_size/2, text=wd, font=("Meiryo UI", 9, "bold"))
            current_day = 1
            for week in range(6):
                for dow in range(7):
                    x = grid_start_x + dow * cell_size; y = grid_start_y + (week+1) * cell_size
                    if week == 0 and dow < first_wd - 1: continue
                    if current_day > last_day: break
                    day_activities = []
                    for label_idx in range(1, len(self.LABELS)-1):
                        for col_idx in range(self.BLOCK_W):
                            ck = (current_day, label_idx, col_idx)
                            if ck in self.cell_data and self.cell_data[ck] == person:
                                label = self.LABELS[label_idx]
                                if label not in day_activities: day_activities.append(label)
                    if "休み" in day_activities:
                        cell_color = "#ff6b6b"; text_color = "white"; border_color = "#cc0000"
                        border_width = 3; font_weight = "bold"; show_yasumi_mark = True
                    elif day_activities:
                        first_activity = day_activities[0]
                        cell_color = self._get_or_assign_color(first_activity)
                        text_color = "black"; border_color = "gray"; border_width = 1
                        font_weight = "normal"; show_yasumi_mark = False
                    else:
                        cell_color = "white"; text_color = "#cccccc"; border_color = "gray"
                        border_width = 1; font_weight = "normal"; show_yasumi_mark = False
                    self.personal_canvas.create_rectangle(x, y, x+cell_size, y+cell_size,
                        fill=cell_color, outline=border_color, width=border_width)
                    if show_yasumi_mark:
                        self.personal_canvas.create_text(x+5, y+5, text=str(current_day),
                            font=("Meiryo UI", 7, font_weight), fill=text_color, anchor="nw")
                        self.personal_canvas.create_text(x+cell_size/2, y+cell_size/2+3, text="休",
                            font=("Meiryo UI", 14, "bold"), fill=text_color)
                    else:
                        self.personal_canvas.create_text(x+cell_size/2, y+cell_size/2, text=str(current_day),
                            font=("Meiryo UI", 8, font_weight), fill=text_color)
                    current_day += 1
            legend_y = grid_start_y + 7 * cell_size + 10
            person_summary = defaultdict(int)
            for (day, label_idx, col_idx), value in self.cell_data.items():
                if value == person:
                    label = self.LABELS[label_idx] if label_idx < len(self.LABELS) else ""
                    if label and label.strip(): person_summary[label] += 1
            legend_text = " | ".join([f"{k}:{v}" for k, v in sorted(person_summary.items())])
            self.personal_canvas.create_text(base_x+cal_width/2, legend_y, text=legend_text,
                font=("Meiryo UI", 8), fill="#666666")
        total_rows = (len(sorted_persons) + cols - 1) // cols
        total_height = start_y + total_rows * (cal_height + margin) + margin
        total_width = start_x + cols * (cal_width + margin)
        self.personal_canvas.configure(scrollregion=(0, 0, total_width, total_height))

    def _bg_color(self, label_idx: int, dow: int) -> str:
        if label_idx == 0 or label_idx == 9 or label_idx == 10: return "#f0f0f0"
        if dow == 5: return "#f0f0ff"
        if dow == 6: return "#fff0f0"
        return "white"

    def _lighten_color(self, color: str, factor: float = 0.6) -> str:
        color = color.lstrip('#')
        if len(color) != 6: return "#e0e0e0"
        try:
            r = int(color[0:2], 16); g = int(color[2:4], 16); b = int(color[4:6], 16)
            r = int(r * factor + 255 * (1 - factor))
            g = int(g * factor + 255 * (1 - factor))
            b = int(b * factor + 255 * (1 - factor))
            return f"#{r:02x}{g:02x}{b:02x}"
        except: return "#e0e0e0"

    def _draw_label_column(self, x, row_y):
        label_fg = {"当直": "#960000", "明け": "#c04040", "休み": "#000096", "ER": "#006400", "外勤": "#640064"}
        for i, lbl in enumerate(self.LABELS):
            y1 = row_y + i * self.CELL_H; y2 = y1 + self.CELL_H
            bg = "#f0f0f0" if (i == 0 or i == 9 or i == 10) else "white"
            self.canvas.create_rectangle(x, y1, x+self.LABEL_COL_W, y2, fill=bg, outline="gray", tags="grid")
            if lbl:
                fg = label_fg.get(lbl, "black")
                self.canvas.create_text(x+self.LABEL_COL_W/2, (y1+y2)/2, text=lbl,
                    font=("Meiryo UI", 10, "bold"), fill=fg, tags="grid")
            if i == 5:
                self.canvas.create_line(x, y2, x+self.LABEL_COL_W, y2, fill="#666666", width=2, dash=(4,4), tags="divider_line")

    def _draw_adjacent_month_block(self, col_x, row_y, day, dow, week, is_prev_month=True):
        self._day_positions[day] = (week, dow, col_x, row_y)
        for li in range(self.BLOCK_H):
            for ci in range(self.BLOCK_W):
                x1 = col_x + ci * self.CELL_W; y1 = row_y + li * self.CELL_H
                x2 = x1 + self.CELL_W; y2 = y1 + self.CELL_H
                base_bg = self._bg_color(li, dow)
                bg = "#f5f5f5" if base_bg == "white" else self._lighten_color(base_bg, 0.8)
                rid = self.canvas.create_rectangle(x1, y1, x2, y2, fill=bg, outline="gray", tags="grid")
                tid = self.canvas.create_text((x1+x2)/2, (y1+y2)/2, text="", font=("Meiryo UI", 9), fill="#888888", tags="cell_text")
                if li == 0 and ci == 0:
                    display_day = abs(day) if day < 0 else day - calendar.monthrange(self.current_year, self.current_month)[1]
                    self.canvas.create_text(x1+3, y1+3, anchor="nw", text=str(display_day),
                        font=("Meiryo UI", 6), fill="#999999", tags="day_number")
                if li == 5 and ci == 0:
                    self.canvas.create_line(col_x, y2, col_x+self.BLOCK_W*self.CELL_W, y2,
                        fill="#666666", width=2, dash=(4,4), tags="divider_line")
                ck = (day, li, ci)
                self._rect_ids[ck] = rid; self._text_ids[ck] = tid; self._cell_bounds[ck] = (x1, y1, x2, y2)

    def _draw_day_block(self, col_x, row_y, day, dow, week):
        self._day_positions[day] = (week, dow, col_x, row_y)
        # ✨ 改善: 祝日判定
        is_holiday = day in self._current_holidays
        for li in range(self.BLOCK_H):
            for ci in range(self.BLOCK_W):
                x1 = col_x + ci * self.CELL_W; y1 = row_y + li * self.CELL_H
                x2 = x1 + self.CELL_W; y2 = y1 + self.CELL_H
                bg = self._bg_color(li, dow)
                # ✨ 祝日セルを薄いオレンジで塗る
                if is_holiday and (li == 0 or bg == "white"):
                    bg = "#fff0e8" if bg == "white" else bg
                rid = self.canvas.create_rectangle(x1, y1, x2, y2, fill=bg, outline="gray", tags="grid")
                tid = self.canvas.create_text((x1+x2)/2, (y1+y2)/2, text="", font=("Meiryo UI", 9), tags="cell_text")
                if li == 0 and ci == 0:
                    # ✨ 改善: 祝日は赤橙色、土は青、日は赤
                    if is_holiday:
                        day_fg = "#cc3300"
                    elif dow == 5:
                        day_fg = "blue"
                    elif dow == 6:
                        day_fg = "red"
                    else:
                        day_fg = "black"
                    self.canvas.create_text(x1+3, y1+3, anchor="nw", text=str(day),
                        font=("Meiryo UI", 6, "bold"), fill=day_fg, tags="day_number")
                    # ✨ 改善: 祝日に「祝」マーク表示
                    if is_holiday:
                        self.canvas.create_text(x2-3, y1+3, anchor="ne", text="祝",
                            font=("Meiryo UI", 5, "bold"), fill="#cc3300", tags="day_number")
                if li == 5 and ci == 0:
                    self.canvas.create_line(col_x, y2, col_x+self.BLOCK_W*self.CELL_W, y2,
                        fill="#666666", width=2, dash=(4,4), tags="divider_line")
                ck = (day, li, ci)
                self._rect_ids[ck] = rid; self._text_ids[ck] = tid; self._cell_bounds[ck] = (x1, y1, x2, y2)


    def _apply_data(self):
        for (day, li, ci), value in self.cell_data.items():
            rid = self._rect_ids.get((day, li, ci)); tid = self._text_ids.get((day, li, ci))
            color = self.cell_colors.get((day, li, ci), "white")
            if rid: self.canvas.itemconfigure(rid, fill=color); self.canvas.tag_raise(rid)
            if tid: self.canvas.itemconfigure(tid, text=value)
        self.canvas.tag_raise("divider_line"); self._raise_all_text(); self._redraw_lock_borders()

    def _update_cell_display(self, cell_key):
        rid = self._rect_ids.get(cell_key); tid = self._text_ids.get(cell_key)
        if cell_key in self.cell_data:
            value = self.cell_data[cell_key]; color = self.cell_colors[cell_key]
            if rid: self.canvas.itemconfigure(rid, fill=color); self.canvas.tag_raise(rid)
            if tid: self.canvas.itemconfigure(tid, text=value); self.canvas.tag_raise(tid)
        else:
            day, li, ci = cell_key; info = self._day_positions.get(day); dow = info[1] if info else 0
            bg = self._bg_color(li, dow)
            if day <= 0 or day > 31:
                bg = "#f5f5f5" if bg == "white" else self._lighten_color(bg, 0.8)
            if rid: self.canvas.itemconfigure(rid, fill=bg); self.canvas.tag_lower(rid)
            if tid: self.canvas.itemconfigure(tid, text=""); self.canvas.tag_raise(tid)

    def _draw_all_lines(self):
        self.canvas.delete("connection_line")
        cells_by_value = defaultdict(list)
        for ck, value in self.cell_data.items():
            cells_by_value[value].append(ck)
        for value, cells in cells_by_value.items():
            sorted_cells = sorted(cells, key=lambda c: (c[0], c[2], c[1]))
            for i in range(len(sorted_cells) - 1):
                ck1 = sorted_cells[i]; ck2 = sorted_cells[i+1]
                day1, li1, _ = ck1; day2, li2, _ = ck2
                if li1 == 0 or li1 == 9 or li1 == 10 or li2 == 0 or li2 == 9 or li2 == 10: continue
                info1 = self._day_positions.get(day1); info2 = self._day_positions.get(day2)
                if info1 is None or info2 is None: continue
                if info1[0] != info2[0]: continue
                bounds1 = self._cell_bounds.get(ck1); bounds2 = self._cell_bounds.get(ck2)
                if bounds1 and bounds2:
                    x1s,y1s,x1e,y1e = bounds1; x2s,y2s,x2e,y2e = bounds2
                    x1c=(x1s+x1e)/2; y1c=(y1s+y1e)/2; x2c=(x2s+x2e)/2; y2c=(y2s+y2e)/2
                    if x2c > x1e: lx1,ly1 = x1e,y1c
                    elif x2c < x1s: lx1,ly1 = x1s,y1c
                    elif y2c > y1e: lx1,ly1 = x1c,y1e
                    elif y2c < y1s: lx1,ly1 = x1c,y1s
                    else: lx1,ly1 = x1c,y1c
                    if x1c > x2e: lx2,ly2 = x2e,y2c
                    elif x1c < x2s: lx2,ly2 = x2s,y2c
                    elif y1c > y2e: lx2,ly2 = x2c,y2e
                    elif y1c < y2s: lx2,ly2 = x2c,y2s
                    else: lx2,ly2 = x2c,y2c
                    line_color = self.cell_colors.get(ck1, "blue")
                    self.canvas.create_line(lx1, ly1, lx2, ly2, fill=line_color, width=2, tags="connection_line")
        self.canvas.tag_raise("connection_line", "grid")
        for ck in self.cell_data.keys():
            rid = self._rect_ids.get(ck)
            if rid: self.canvas.tag_raise(rid)
        self.canvas.tag_raise("divider_line"); self._raise_all_text(); self._redraw_lock_borders()

    def _raise_all_text(self):
        for ck, tid in self._text_ids.items():
            if tid: self.canvas.tag_raise(tid)

    def _cell_at(self, cx, cy) -> Optional[CellKey]:
        for ck, (x1, y1, x2, y2) in self._cell_bounds.items():
            if x1 <= cx <= x2 and y1 <= cy <= y2: return ck
        return None

    def _on_click(self, event):
        self._finish_edit()
        cx = self.canvas.canvasx(event.x); cy = self.canvas.canvasy(event.y)
        ck = self._cell_at(cx, cy)
        if ck is not None and ck in self.cell_data:
            self._drag_source = ck; self._drag_start_cx = cx; self._drag_start_cy = cy; self._dragging = False

    def _on_motion(self, event):
        if self._drag_source is None: return
        if self._drag_source in self.locked_cells: return
        cx = self.canvas.canvasx(event.x); cy = self.canvas.canvasy(event.y)
        dx = abs(cx - self._drag_start_cx); dy = abs(cy - self._drag_start_cy)
        if not self._dragging and (dx > self.DRAG_THRESHOLD or dy > self.DRAG_THRESHOLD):
            self._dragging = True; self._show_drag_ghost(self._drag_source, cx, cy)
        if self._dragging:
            self._update_drag_ghost(cx, cy)
            self._update_drop_highlight(self._cell_at(cx, cy))

    def _on_release(self, event):
        if not self._dragging:
            if self._drag_source is not None: self._toggle_selection(self._drag_source)
            self._drag_source = None; return
        if self._drag_source in self.locked_cells:
            self._drag_source = None; self._dragging = False; self._clear_drag_visuals(); return
        cx = self.canvas.canvasx(event.x); cy = self.canvas.canvasy(event.y)
        drop_ck = self._cell_at(cx, cy)
        if drop_ck and drop_ck in self.locked_cells:
            self._drag_source = None; self._dragging = False; self._clear_drag_visuals()
            messagebox.showwarning("ロック中", "このセルはロックされているため移動できません"); return
        if drop_ck and self._drag_source:
            self._save_state()
            self._remove_special_border(self._drag_source); self._remove_special_border(drop_ck)
            src_val = self.cell_data.get(self._drag_source, ""); src_color = self.cell_colors.get(self._drag_source, "white")
            src_seq = self.cell_sequence.get(self._drag_source, None); src_er = self._drag_source in self.er_marks
            dst_val = self.cell_data.get(drop_ck, ""); dst_color = self.cell_colors.get(drop_ck, "white")
            dst_seq = self.cell_sequence.get(drop_ck, None); dst_er = drop_ck in self.er_marks
            if dst_val:
                self.cell_data[self._drag_source] = dst_val; self.cell_colors[self._drag_source] = dst_color
                if dst_seq is not None: self.cell_sequence[self._drag_source] = dst_seq
                else: self.cell_sequence.pop(self._drag_source, None)
                if dst_er: self.er_marks.add(self._drag_source)
                else: self.er_marks.discard(self._drag_source)
            else:
                self.cell_data.pop(self._drag_source, None); self.cell_colors.pop(self._drag_source, None)
                self.cell_sequence.pop(self._drag_source, None); self.er_marks.discard(self._drag_source)
            if src_val:
                self.cell_data[drop_ck] = src_val; self.cell_colors[drop_ck] = src_color
                if src_seq is not None: self.cell_sequence[drop_ck] = src_seq
                if src_er: self.er_marks.add(drop_ck)
            else:
                self.cell_data.pop(drop_ck, None); self.cell_colors.pop(drop_ck, None)
                self.cell_sequence.pop(drop_ck, None); self.er_marks.discard(drop_ck)
            self._update_cell_display(self._drag_source); self._update_cell_display(drop_ck)
            self._draw_all_lines()
            _, drop_li, _ = drop_ck
            drop_label = self.LABELS[drop_li] if drop_li < len(self.LABELS) else ""
            if drop_label and drop_label.strip(): self.sound_manager.play_for_label(drop_label)
            else: self.sound_manager.play("move")
            self.draw_statistics(); self._clear_selection()
        self._drag_source = None; self._dragging = False; self._clear_drag_visuals()

    def _show_drag_ghost(self, ck, cx, cy):
        val = self.cell_data.get(ck, ""); color = self.cell_colors.get(ck, "white")
        bounds = self._cell_bounds.get(ck)
        if not bounds: return
        x1, y1, x2, y2 = bounds; w = x2-x1; h = y2-y1
        self._drag_ghost_rect = self.canvas.create_rectangle(cx-w/2, cy-h/2, cx+w/2, cy+h/2,
            fill=color, outline="blue", width=2, stipple="gray50", tags="drag_ghost")
        self._drag_ghost_text = self.canvas.create_text(cx, cy, text=val, font=("Meiryo UI", 9), tags="drag_ghost")

    def _update_drag_ghost(self, cx, cy):
        if self._drag_ghost_rect:
            bounds = self._cell_bounds.get(self._drag_source, (0,0,60,25))
            x1,y1,x2,y2 = bounds; w=x2-x1; h=y2-y1
            self.canvas.coords(self._drag_ghost_rect, cx-w/2, cy-h/2, cx+w/2, cy+h/2)
        if self._drag_ghost_text: self.canvas.coords(self._drag_ghost_text, cx, cy)

    def _update_drop_highlight(self, drop_ck):
        if self._drag_highlight: self.canvas.delete(self._drag_highlight); self._drag_highlight = None
        if drop_ck and drop_ck != self._drag_source:
            bounds = self._cell_bounds.get(drop_ck)
            if bounds:
                x1,y1,x2,y2 = bounds
                self._drag_highlight = self.canvas.create_rectangle(x1,y1,x2,y2, outline="red", width=3, tags="drag_highlight")

    def _clear_drag_visuals(self):
        self.canvas.delete("drag_ghost"); self.canvas.delete("drag_highlight")
        self._drag_ghost_rect = None; self._drag_ghost_text = None; self._drag_highlight = None

    def _toggle_selection(self, cell_key):
        if cell_key not in self.cell_data:
            if self.selected_value is not None: self._clear_selection()
            return
        _, li, _ = cell_key
        if li == 0 or li == 9 or li == 10: return
        clicked_value = self.cell_data[cell_key]
        if self.selected_value == clicked_value: self._clear_selection()
        else: self.selected_value = clicked_value; self._apply_selection_mask()

    def _apply_selection_mask(self):
        if self.selected_value is None: return
        for ck, value in self.cell_data.items():
            rid = self._rect_ids.get(ck); tid = self._text_ids.get(ck)
            if value != self.selected_value:
                original_color = self.cell_colors.get(ck, "white")
                masked_color = self._lighten_color(original_color, self.mask_opacity)
                if rid: self.canvas.itemconfigure(rid, fill=masked_color)
                if tid: self.canvas.itemconfigure(tid, fill="#cccccc")
        self._apply_line_mask()

    def _apply_line_mask(self):
        if self.selected_value is None: return
        self.canvas.delete("connection_line"); self.canvas.delete("connection_line_selected")
        cells_by_value = defaultdict(list)
        for ck, value in self.cell_data.items(): cells_by_value[value].append(ck)
        for value, cells in cells_by_value.items():
            if value == self.selected_value: continue
            sorted_cells = sorted(cells, key=lambda c: (c[0], c[2], c[1]))
            original_color = self.cell_colors.get(sorted_cells[0], "blue")
            line_color = self._lighten_color(original_color, self.mask_opacity)
            for i in range(len(sorted_cells)-1):
                ck1=sorted_cells[i]; ck2=sorted_cells[i+1]
                day1,li1,_=ck1; day2,li2,_=ck2
                if li1==0 or li1==9 or li1==10 or li2==0 or li2==9 or li2==10: continue
                info1=self._day_positions.get(day1); info2=self._day_positions.get(day2)
                if info1 is None or info2 is None: continue
                if info1[0] != info2[0]: continue
                bounds1=self._cell_bounds.get(ck1); bounds2=self._cell_bounds.get(ck2)
                if bounds1 and bounds2:
                    x1s,y1s,x1e,y1e=bounds1; x2s,y2s,x2e,y2e=bounds2
                    x1c=(x1s+x1e)/2; y1c=(y1s+y1e)/2; x2c=(x2s+x2e)/2; y2c=(y2s+y2e)/2
                    if x2c>x1e: lx1,ly1=x1e,y1c
                    elif x2c<x1s: lx1,ly1=x1s,y1c
                    elif y2c>y1e: lx1,ly1=x1c,y1e
                    elif y2c<y1s: lx1,ly1=x1c,y1s
                    else: lx1,ly1=x1c,y1c
                    if x1c>x2e: lx2,ly2=x2e,y2c
                    elif x1c<x2s: lx2,ly2=x2s,y2c
                    elif y1c>y2e: lx2,ly2=x2c,y2e
                    elif y1c<y2s: lx2,ly2=x2c,y2s
                    else: lx2,ly2=x2c,y2c
                    self.canvas.create_line(lx1,ly1,lx2,ly2, fill=line_color, width=2, tags="connection_line")
        self.canvas.tag_raise("connection_line", "grid")
        for ck in self.cell_data.keys():
            if self.cell_data.get(ck) != self.selected_value:
                rid = self._rect_ids.get(ck)
                if rid: self.canvas.tag_raise(rid)
        if self.selected_value in cells_by_value:
            cells = cells_by_value[self.selected_value]
            sorted_cells = sorted(cells, key=lambda c: (c[0], c[2], c[1]))
            line_color = self.cell_colors.get(sorted_cells[0], "blue")
            for i in range(len(sorted_cells)-1):
                ck1=sorted_cells[i]; ck2=sorted_cells[i+1]
                day1,li1,_=ck1; day2,li2,_=ck2
                if li1==0 or li1==9 or li1==10 or li2==0 or li2==9 or li2==10: continue
                info1=self._day_positions.get(day1); info2=self._day_positions.get(day2)
                if info1 is None or info2 is None: continue
                if info1[0] != info2[0]: continue
                bounds1=self._cell_bounds.get(ck1); bounds2=self._cell_bounds.get(ck2)
                if bounds1 and bounds2:
                    x1s,y1s,x1e,y1e=bounds1; x2s,y2s,x2e,y2e=bounds2
                    x1c=(x1s+x1e)/2; y1c=(y1s+y1e)/2; x2c=(x2s+x2e)/2; y2c=(y2s+y2e)/2
                    if x2c>x1e: lx1,ly1=x1e,y1c
                    elif x2c<x1s: lx1,ly1=x1s,y1c
                    elif y2c>y1e: lx1,ly1=x1c,y1e
                    elif y2c<y1s: lx1,ly1=x1c,y1s
                    else: lx1,ly1=x1c,y1c
                    if x1c>x2e: lx2,ly2=x2e,y2c
                    elif x1c<x2s: lx2,ly2=x2s,y2c
                    elif y1c>y2e: lx2,ly2=x2c,y2e
                    elif y1c<y2s: lx2,ly2=x2c,y2s
                    else: lx2,ly2=x2c,y2c
                    self.canvas.create_line(lx1,ly1,lx2,ly2, fill=line_color, width=3, tags="connection_line_selected")
            for ck in cells:
                rid = self._rect_ids.get(ck)
                if rid: self.canvas.tag_raise(rid)
        self.canvas.tag_raise("divider_line"); self.canvas.tag_raise("connection_line_selected")
        self._raise_all_text(); self._redraw_lock_borders()

    def _clear_selection(self):
        self.selected_value = None
        self.canvas.delete("connection_line_selected")
        for ck, value in self.cell_data.items():
            rid = self._rect_ids.get(ck); tid = self._text_ids.get(ck)
            original_color = self.cell_colors.get(ck, "white")
            if rid: self.canvas.itemconfigure(rid, fill=original_color)
            if tid: self.canvas.itemconfigure(tid, fill="black")
        self._draw_all_lines()

    def _add_lock(self, cell_key):
        if cell_key not in self.locked_cells:
            self.locked_cells.add(cell_key); self._draw_lock_border(cell_key)
            messagebox.showinfo("ロック", "このセルをロックしました")

    def _remove_lock(self, cell_key):
        if cell_key in self.locked_cells:
            self.locked_cells.discard(cell_key)
            if cell_key in self.locked_cell_borders:
                self.canvas.delete(self.locked_cell_borders[cell_key])
                del self.locked_cell_borders[cell_key]
            messagebox.showinfo("ロック解除", "このセルのロックを解除しました")

    def _draw_lock_border(self, cell_key):
        bounds = self._cell_bounds.get(cell_key)
        if not bounds: return
        x1,y1,x2,y2 = bounds
        border_id = self.canvas.create_rectangle(x1,y1,x2,y2, outline="black", width=3, tags="lock_border")
        self.locked_cell_borders[cell_key] = border_id
        self.canvas.tag_raise("lock_border"); self._raise_all_text()

    def _redraw_lock_borders(self):
        for border_id in self.locked_cell_borders.values(): self.canvas.delete(border_id)
        self.locked_cell_borders.clear()
        for cell_key in list(self.locked_cells):
            if cell_key not in self.cell_data: self.locked_cells.discard(cell_key); continue
            self._draw_lock_border(cell_key)
        self.canvas.tag_raise("lock_border"); self.canvas.tag_raise("special_border"); self._raise_all_text()

    def _show_gray_cell_menu(self, event, ck):
        ICONS = {"外勤":"🚗","委員会":"🏛","コース":"📚","訓練":"🎯"}
        popup = tk.Menu(self.root, tearoff=0)
        popup.add_command(label="🚑 ドクターカー", command=lambda c=ck: self._set_gray_cell_text(c, "D/C"))
        popup.add_separator()
        for cat in ["外勤","委員会","コース","訓練"]:
            icon = ICONS[cat]; sub = tk.Menu(popup, tearoff=0)
            items = self.special_items[cat]
            if items:
                for item in items:
                    sub.add_command(label=item, command=lambda c=ck, t=item: self._set_gray_cell_text(c, t))
                sub.add_separator()
                sub.add_command(label="（クリア）", command=lambda c=ck: self._set_gray_cell_text(c, ""))
            else:
                sub.add_command(label=f"（{cat}の項目未登録）", state=tk.DISABLED)
                sub.add_command(label=f"→ {cat}を登録する...", command=lambda c=cat: self._open_special_item_window(c))
            popup.add_cascade(label=f"{icon} {cat}", menu=sub)
        current_val = self.cell_data.get(ck, "")
        if current_val.strip():
            popup.add_separator()
            popup.add_command(label="✖ このセルをクリア", command=lambda c=ck: self._set_gray_cell_text(c, ""))
        try: popup.tk_popup(event.x_root, event.y_root)
        finally: popup.grab_release()

    def _set_gray_cell_text(self, ck, text: str):
        self._save_state()
        if text: self.cell_data[ck] = text; self.cell_colors[ck] = "#d8d8d8"
        else: self.cell_data.pop(ck, None); self.cell_colors.pop(ck, None)
        self._update_cell_display(ck); self.draw_statistics()

    def _on_right_click(self, event):
        cx = self.canvas.canvasx(event.x); cy = self.canvas.canvasy(event.y)
        ck = self._cell_at(cx, cy)
        if ck is None: return
        day, li, ci = ck
        if li == 0 or li == 9 or li == 10:
            self._show_gray_cell_menu(event, ck); return
        popup = tk.Menu(self.root, tearoff=0)
        if ck in self.cell_data:
            if ck in self.locked_cells:
                popup.add_command(label="🔓 ロック解除", command=lambda: self._remove_lock(ck))
            else:
                popup.add_command(label="🔒 ロック", command=lambda: self._add_lock(ck))
            popup.add_separator()
            current_border = self.special_borders.get(ck, '')
            dc_check = "✔ " if current_border in ('doctor_car','both') else "　"
            ang_check = "✔ " if current_border in ('angio_call','both') else "　"
            popup.add_command(label=f"{dc_check}ドクターカー", command=lambda: self._toggle_special_border(ck, 'doctor_car'))
            popup.add_command(label=f"{ang_check}アンギオコール", command=lambda: self._toggle_special_border(ck, 'angio_call'))
            popup.add_separator()
            popup.add_command(label="枠線を解除", command=lambda: self._remove_special_border(ck))
            popup.add_separator()
        team_defs = [("A","🔵 Aチーム","#1d4ed8"),("B","🟢 Bチーム","#15803d"),("非常勤","⬜ 非常勤","#78716c")]
        for team_key, team_label, team_color in team_defs:
            team_members = [m for m in self.members_data if m.get("active",True) and m.get("team","A")==team_key]
            sub = tk.Menu(popup, tearoff=0)
            if team_members:
                for m in team_members:
                    dname = m["display_name"]
                    sub.add_command(label=dname,
                        command=lambda k=ck, d=dname, c=m.get("color","#cccccc"): self._set_cell_from_menu(k, d, c))
            else:
                sub.add_command(label="（メンバーなし）", state=tk.DISABLED)
            popup.add_cascade(label=team_label, menu=sub)
        if ck in self.cell_data and self.cell_data[ck].strip():
            popup.add_separator()
            popup.add_command(label="✖ クリア", command=lambda k=ck: self._clear_cell_from_menu(k))
        try: popup.tk_popup(event.x_root, event.y_root)
        finally: popup.grab_release()

    def _set_cell_from_menu(self, ck, member_name: str, color: str):
        if ck in self.locked_cells: messagebox.showwarning("ロック中", "このセルはロックされています"); return
        self._save_state()
        self.cell_data[ck] = member_name; self.cell_colors[ck] = color
        if ck not in self.cell_sequence: self.cell_sequence[ck] = self.next_sequence; self.next_sequence += 1
        self._update_cell_display(ck); self._draw_all_lines()
        _, li, _ = ck
        lbl = self.LABELS[li] if li < len(self.LABELS) else ""
        if lbl and lbl.strip(): self.sound_manager.play_for_label(lbl)
        self.draw_statistics()

    def _clear_cell_from_menu(self, ck):
        if ck in self.locked_cells: messagebox.showwarning("ロック中", "このセルはロックされています"); return
        self._save_state()
        self.cell_data.pop(ck, None); self.cell_colors.pop(ck, None)
        self.cell_sequence.pop(ck, None); self.er_marks.discard(ck)
        self._update_cell_display(ck); self._draw_all_lines(); self.draw_statistics()

    def _draw_border_visual(self, cell_key, border_type):
        bounds = self._cell_bounds.get(cell_key)
        if not bounds: return []
        x1,y1,x2,y2 = bounds; ids = []; PINK="#FF1493"; BLUE="#00BFFF"; W=4
        if border_type == 'doctor_car':
            ids.append(self.canvas.create_rectangle(x1,y1,x2,y2, outline=PINK, width=W, tags="special_border"))
        elif border_type == 'angio_call':
            ids.append(self.canvas.create_rectangle(x1,y1,x2,y2, outline=BLUE, width=W, tags="special_border"))
        elif border_type == 'both':
            ids.append(self.canvas.create_line(x1,y1,x2,y1, fill=PINK, width=W, tags="special_border"))
            ids.append(self.canvas.create_line(x1,y1,x1,y2, fill=PINK, width=W, tags="special_border"))
            ids.append(self.canvas.create_line(x1,y2,x2,y2, fill=BLUE, width=W, tags="special_border"))
            ids.append(self.canvas.create_line(x2,y1,x2,y2, fill=BLUE, width=W, tags="special_border"))
        return ids

    def _toggle_special_border(self, cell_key, border_type):
        current = self.special_borders.get(cell_key, '')
        if current == '': new_type = border_type
        elif current == border_type: self._remove_special_border(cell_key); return
        elif current == 'both': new_type = 'angio_call' if border_type == 'doctor_car' else 'doctor_car'
        else: new_type = 'both'
        self._remove_special_border(cell_key)
        ids = self._draw_border_visual(cell_key, new_type)
        self.special_borders[cell_key] = new_type; self.special_border_ids[cell_key] = ids
        self.canvas.tag_raise("special_border"); self._raise_all_text()
        if new_type == 'doctor_car': self.sound_manager.play("vroom")
        elif new_type == 'angio_call': self.sound_manager.play("sheen")
        elif new_type == 'both': self.sound_manager.play("vroom")

    def _add_special_border(self, cell_key, border_type):
        self._remove_special_border(cell_key)
        ids = self._draw_border_visual(cell_key, border_type)
        self.special_borders[cell_key] = border_type; self.special_border_ids[cell_key] = ids
        self.canvas.tag_raise("special_border"); self._raise_all_text()
        if border_type == 'doctor_car': self.sound_manager.play("vroom")
        elif border_type == 'angio_call': self.sound_manager.play("sheen")

    def _remove_special_border(self, cell_key):
        if cell_key in self.special_border_ids:
            for item_id in self.special_border_ids[cell_key]: self.canvas.delete(item_id)
            del self.special_border_ids[cell_key]
        if cell_key in self.special_borders: del self.special_borders[cell_key]

    def _redraw_special_borders(self):
        for ids in self.special_border_ids.values():
            for item_id in ids: self.canvas.delete(item_id)
        self.special_border_ids.clear()
        for cell_key, border_type in list(self.special_borders.items()):
            if cell_key not in self.cell_data: del self.special_borders[cell_key]; continue
            if not self._cell_bounds.get(cell_key): continue
            ids = self._draw_border_visual(cell_key, border_type)
            self.special_border_ids[cell_key] = ids
        self.canvas.tag_raise("special_border"); self._raise_all_text()

    def _on_double_click(self, event):
        self._drag_source = None; self._dragging = False; self._clear_drag_visuals(); self._clear_selection()
        cx = self.canvas.canvasx(event.x); cy = self.canvas.canvasy(event.y)
        ck = self._cell_at(cx, cy)
        if ck is None: return
        self._finish_edit(); self._start_edit(ck)

    def _start_edit(self, cell_key):
        bounds = self._cell_bounds.get(cell_key)
        if bounds is None: return
        x1,y1,x2,y2 = bounds
        current_val = self.cell_data.get(cell_key, "")
        color = self.cell_colors.get(cell_key, "white")
        self._edit_entry = tk.Entry(self.canvas, font=("Meiryo UI", 9), justify=tk.CENTER, bg=color)
        self._edit_entry.insert(0, current_val); self._edit_entry.select_range(0, tk.END)
        self._edit_entry.focus_set()
        self._edit_entry.bind("<Return>", lambda e: self._finish_edit())
        self._edit_entry.bind("<Escape>", lambda e: self._cancel_edit())
        self._edit_entry.bind("<Tab>", lambda e: self._tab_to_next(e))
        self._edit_window_id = self.canvas.create_window((x1+x2)/2, (y1+y2)/2,
            window=self._edit_entry, width=x2-x1-2, height=y2-y1-2, tags="edit_entry")
        self._edit_key = cell_key

    def _tab_to_next(self, event):
        if self._edit_key is None: return "break"
        day, li, ci = self._edit_key; self._finish_edit()
        next_key = (day, li, ci+1)
        if next_key not in self._cell_bounds: next_key = (day, li+1, 0)
        if next_key in self._cell_bounds: self._start_edit(next_key)
        return "break"

    def _finish_edit(self):
        if self._edit_entry is None: return
        new_val = self._edit_entry.get().strip(); ck = self._edit_key
        self.canvas.delete("edit_entry"); self._edit_entry.destroy()
        self._edit_entry = None; self._edit_key = None
        if ck is None: return
        old_val = self.cell_data.get(ck, "")
        if new_val == old_val: return
        self._save_state()
        rid = self._rect_ids.get(ck); tid = self._text_ids.get(ck)
        if not new_val:
            self.cell_data.pop(ck, None); self.cell_colors.pop(ck, None)
            self.cell_sequence.pop(ck, None); self.er_marks.discard(ck)
            if rid:
                day,li,ci = ck; info = self._day_positions.get(day); dow = info[1] if info else 0
                bg = self._bg_color(li, dow)
                if day <= 0 or day > 31:
                    bg = "#f5f5f5" if bg == "white" else self._lighten_color(bg, 0.8)
                self.canvas.itemconfigure(rid, fill=bg)
            if tid: self.canvas.itemconfigure(tid, text="")
        else:
            self.cell_data[ck] = new_val
            if ck not in self.cell_sequence: self.cell_sequence[ck] = self.next_sequence; self.next_sequence += 1
            color = self._get_or_assign_color(new_val); self.cell_colors[ck] = color
            if rid: self.canvas.itemconfigure(rid, fill=color)
            if tid: self.canvas.itemconfigure(tid, text=new_val)
            day, li, ci = ck
            if day == 1 and li >= 1 and li <= 7: self._auto_replicate(new_val, color, li, ci)
        self._draw_all_lines(); self.draw_statistics(); self._clear_selection()

    def _auto_replicate(self, value, color, label_idx, col_idx):
        last_day = calendar.monthrange(self.current_year, self.current_month)[1]
        for target_day in range(2, last_day+1):
            target_key = (target_day, label_idx, col_idx)
            if target_key in self.cell_data: continue
            self.cell_data[target_key] = value; self.cell_colors[target_key] = color
            if target_key not in self.cell_sequence:
                self.cell_sequence[target_key] = self.next_sequence; self.next_sequence += 1
            self._update_cell_display(target_key)

    def _cancel_edit(self):
        if self._edit_entry:
            self.canvas.delete("edit_entry"); self._edit_entry.destroy()
            self._edit_entry = None; self._edit_key = None

    def _get_or_assign_color(self, value):
        if value in self.value_colors: return self.value_colors[value]
        import hashlib
        h = int(hashlib.md5(value.encode()).hexdigest()[:6], 16)
        hue = (h % 360) / 360.0; s, l = 0.45, 0.88
        def hue2rgb(p, q, t):
            if t < 0: t += 1
            if t > 1: t -= 1
            if t < 1/6: return p + (q-p) * 6 * t
            if t < 1/2: return q
            if t < 2/3: return p + (q-p) * (2/3-t) * 6
            return p
        q = l + s - l*s; p = 2*l - q
        rv = hue2rgb(p, q, hue+1/3); gv = hue2rgb(p, q, hue); bv = hue2rgb(p, q, hue-1/3)
        r, g, b = int(rv*255), int(gv*255), int(bv*255)
        color = f"#{r:02x}{g:02x}{b:02x}"; self.value_colors[value] = color; return color

    def _save_state(self):
        state = {
            'cell_data': copy.deepcopy(self.cell_data), 'cell_colors': copy.deepcopy(self.cell_colors),
            'cell_sequence': copy.deepcopy(self.cell_sequence), 'er_marks': copy.deepcopy(self.er_marks),
            'value_colors': copy.deepcopy(self.value_colors), 'next_sequence': self.next_sequence,
            'locked_cells': copy.deepcopy(self.locked_cells), 'completed_days': copy.deepcopy(self.completed_days),
        }
        self.undo_stack.append(state)
        if len(self.undo_stack) > self.max_history: self.undo_stack.pop(0)
        self.redo_stack.clear()

    def undo(self):
        if not self.undo_stack: messagebox.showinfo("情報", "これ以上戻せません"); return
        current_state = {
            'cell_data': copy.deepcopy(self.cell_data), 'cell_colors': copy.deepcopy(self.cell_colors),
            'cell_sequence': copy.deepcopy(self.cell_sequence), 'er_marks': copy.deepcopy(self.er_marks),
            'value_colors': copy.deepcopy(self.value_colors), 'next_sequence': self.next_sequence,
            'locked_cells': copy.deepcopy(self.locked_cells), 'completed_days': copy.deepcopy(self.completed_days),
        }
        self.redo_stack.append(current_state)
        previous_state = self.undo_stack.pop()
        self.cell_data = previous_state['cell_data']; self.cell_colors = previous_state['cell_colors']
        self.cell_sequence = previous_state['cell_sequence']; self.er_marks = previous_state['er_marks']
        self.value_colors = previous_state['value_colors']; self.next_sequence = previous_state['next_sequence']
        self.locked_cells = previous_state.get('locked_cells', set())
        self.completed_days = previous_state.get('completed_days', set())
        self._redraw_all_cells(); self.draw_statistics(); self._clear_selection()

    def redo(self):
        if not self.redo_stack: messagebox.showinfo("情報", "これ以上進めません"); return
        current_state = {
            'cell_data': copy.deepcopy(self.cell_data), 'cell_colors': copy.deepcopy(self.cell_colors),
            'cell_sequence': copy.deepcopy(self.cell_sequence), 'er_marks': copy.deepcopy(self.er_marks),
            'value_colors': copy.deepcopy(self.value_colors), 'next_sequence': self.next_sequence,
            'locked_cells': copy.deepcopy(self.locked_cells), 'completed_days': copy.deepcopy(self.completed_days),
        }
        self.undo_stack.append(current_state)
        next_state = self.redo_stack.pop()
        self.cell_data = next_state['cell_data']; self.cell_colors = next_state['cell_colors']
        self.cell_sequence = next_state['cell_sequence']; self.er_marks = next_state['er_marks']
        self.value_colors = next_state['value_colors']; self.next_sequence = next_state['next_sequence']
        self.locked_cells = next_state.get('locked_cells', set())
        self.completed_days = next_state.get('completed_days', set())
        self._redraw_all_cells(); self.draw_statistics(); self._clear_selection()

    def _redraw_all_cells(self):
        for ck in self._rect_ids.keys():
            rid = self._rect_ids.get(ck); tid = self._text_ids.get(ck)
            day, li, ci = ck; info = self._day_positions.get(day)
            if info:
                dow = info[1]; bg = self._bg_color(li, dow)
                if day <= 0 or day > 31:
                    bg = "#f5f5f5" if bg == "white" else self._lighten_color(bg, 0.8)
                if rid: self.canvas.itemconfigure(rid, fill=bg); self.canvas.tag_lower(rid)
                if tid: self.canvas.itemconfigure(tid, text="")
        for ck, value in self.cell_data.items():
            rid = self._rect_ids.get(ck); tid = self._text_ids.get(ck)
            color = self.cell_colors.get(ck, "#ffffff")
            if rid: self.canvas.itemconfigure(rid, fill=color); self.canvas.tag_raise(rid)
            if tid: self.canvas.itemconfigure(tid, text=value)
        self._draw_all_lines(); self._redraw_lock_borders(); self._redraw_day_complete_markers()

    def clear_data(self):
        if messagebox.askyesno("確認", "すべてのデータをクリアしますか？"):
            self.cell_data.clear(); self.cell_colors.clear(); self.cell_sequence.clear(); self.er_marks.clear()
            self.value_colors = {"ER":"#98fb98","A":"#ffcc99","B":"#add8e6","当直":"#ffb6c1",
                "外勤":"#dda0dd","出張":"#f0e68c","休み":"#e0ffff","秋":"#FFE4B2","桜":"#FFD1DC",
                "小":"#D8D8E8","金":"#C5EDD6","坪":"#FFDCB5","東":"#DDD0F5","長":"#BDE3F8","矢":"#FFF0B5","宮":"#E0E0E0"}
            self.next_sequence = 1; self.undo_stack.clear(); self.redo_stack.clear()
            self.special_borders.clear(); self.special_border_ids.clear()
            self.locked_cells.clear(); self.locked_cell_borders.clear()
            self.completed_days.clear(); self.day_complete_rects.clear()
            self.create_calendar(); self.draw_statistics()

    def save_data(self):
        filename = filedialog.asksaveasfilename(defaultextension=".json",
            filetypes=[("JSON files","*.json"),("All files","*.*")],
            initialfile=f"calendar_{self.current_year}_{self.current_month}.json")
        if not filename: return
        data = {
            "year": self.current_year, "month": self.current_month,
            "cell_data": {f"{d},{i},{j}": v for (d,i,j),v in self.cell_data.items()},
            "cell_colors": {f"{d},{i},{j}": v for (d,i,j),v in self.cell_colors.items()},
            "cell_sequence": {f"{d},{i},{j}": v for (d,i,j),v in self.cell_sequence.items()},
            "er_marks": [f"{d},{i},{j}" for d,i,j in self.er_marks],
            "value_colors": self.value_colors, "next_sequence": self.next_sequence,
            "special_borders": {f"{d},{i},{j}": v for (d,i,j),v in self.special_borders.items()},
            "locked_cells": [f"{d},{i},{j}" for d,i,j in self.locked_cells],
            "completed_days": list(self.completed_days),
            "team_rules": self.team_rules, "members_data": self.members_data,
            "next_member_id": self._next_member_id,
            "day_memos": {str(k): v for k, v in self.day_memos.items()},
            "special_items": self.special_items,
            "sub_cell_data": {f"{name}|||{d}|||{slot}": v for (name,d,slot),v in self.sub_cell_data.items()},
        }
        with open(filename, "w", encoding="utf-8") as f: json.dump(data, f, ensure_ascii=False, indent=2)
        messagebox.showinfo("保存完了", f"データを保存しました:\n{filename}")

    def load_data(self):
        filename = filedialog.askopenfilename(filetypes=[("JSON files","*.json"),("All files","*.*")])
        if not filename: return
        try:
            with open(filename, "r", encoding="utf-8") as f: data = json.load(f)
            self.current_year = data["year"]; self.current_month = data["month"]
            self.year_var.set(str(self.current_year)); self.month_var.set(str(self.current_month))
            self.cell_data = {tuple(map(int, k.split(","))): v for k, v in data["cell_data"].items()}
            self.cell_colors = {tuple(map(int, k.split(","))): v for k, v in data["cell_colors"].items()}
            self.cell_sequence = {tuple(map(int, k.split(","))): v for k, v in data["cell_sequence"].items()}
            self.er_marks = {tuple(map(int, k.split(","))) for k in data["er_marks"]}
            self.value_colors = data["value_colors"]; self.next_sequence = data["next_sequence"]
            self.special_borders = {tuple(map(int, k.split(","))): v for k, v in data.get("special_borders", {}).items()}
            self.locked_cells = {tuple(map(int, k.split(","))) for k in data.get("locked_cells", [])}
            self.completed_days = set(data.get("completed_days", []))
            self.team_rules = data.get("team_rules", "")
            self.members_data = data.get("members_data", [])
            if not self.members_data: self._load_default_members()
            self._next_member_id = data.get("next_member_id", len(self.members_data)+1)
            if "special_items" in data:
                for cat in ["外勤","委員会","コース","訓練"]:
                    if cat in data["special_items"]: self.special_items[cat] = data["special_items"][cat]
            self.day_memos = {int(k): v for k, v in data.get("day_memos", {}).items()}
            if "sub_cell_data" in data:
                self.sub_cell_data = {}
                for k, v in data["sub_cell_data"].items():
                    parts = k.split("|||")
                    if len(parts) == 3: self.sub_cell_data[(parts[0], int(parts[1]), parts[2])] = v
                    elif len(parts) == 2: self.sub_cell_data[(parts[0], int(parts[1]), "duty")] = v
            else: self.sub_cell_data = {}
            self.undo_stack.clear(); self.redo_stack.clear()
            self.create_calendar(); self.draw_statistics()
            messagebox.showinfo("読込完了", "データを読み込みました")
        except Exception as e: messagebox.showerror("エラー", f"データの読み込みに失敗しました:\n{e}")

    def export_to_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            messagebox.showerror("エラー", "openpyxlがインストールされていません。\npip install openpyxlを実行してください。"); return
        filename = filedialog.asksaveasfilename(defaultextension=".xlsx",
            filetypes=[("Excel files","*.xlsx"),("All files","*.*")],
            initialfile=f"calendar_{self.current_year}_{self.current_month}.xlsx")
        if not filename: return
        try:
            wb = Workbook(); ws = wb.active; ws.title = f"{self.current_year}年{self.current_month}月"
            ws.merge_cells("B1:AE1"); ws["B1"] = f"{self.current_year}年 {self.current_month}月"
            ws["B1"].font = Font(name="Meiryo UI", size=24, bold=True)
            ws["B1"].alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 35
            for (day, li, ci), value in self.cell_data.items():
                info = self._day_positions.get(day)
                if info is None: continue
                week, dow, _, _ = info
                excel_row = 3 + week * self.BLOCK_H + li
                excel_col = 2 + dow * self.BLOCK_W + ci + 1
                cell = ws.cell(row=excel_row, column=excel_col, value=value)
                ck = (day, li, ci)
                if ck in self.cell_colors:
                    h = self.cell_colors[ck].replace("#", "")
                    cell.fill = PatternFill(start_color=h, end_color=h, fill_type="solid")
            wb.save(filename)
            messagebox.showinfo("エクスポート完了", f"Excelファイルを保存しました:\n{filename}")
        except Exception as e: messagebox.showerror("エラー", f"Excelファイルの保存に失敗しました:\n{e}")

    def _load_default_members(self):
        default_names = ["秋","桜","小","金","坪","東","長","矢","宮"]
        for name in default_names:
            if name in self.value_colors:
                self.members_data.append({
                    "id": str(self._next_member_id).zfill(3), "name": name, "display_name": name,
                    "active": True, "skill_level": "専門医", "personal_request": "",
                    "color": self.value_colors[name], "team": "A"
                })
                self._next_member_id += 1

    def open_member_management(self):
        MemberManagementDialog(self)

    def _format_members_for_ai(self):
        active_members = [m for m in self.members_data if m["active"]]
        if not active_members: return "アクティブなメンバーはいません"
        lines = []
        for team in ("A","B","非常勤"):
            label = f"【{team}チーム】" if team in ("A","B") else "【非常勤】"
            team_members = [m for m in active_members if m.get("team","A") == team]
            if team_members:
                lines.append(label)
                for i, member in enumerate(team_members, 1):
                    lines.append(f"  {i}. {member['name']}（{member['skill_level']}）")
                    if member['personal_request']: lines.append(f"     リクエスト: {member['personal_request']}")
        return "\n".join(lines)


    # ── Sub-calendar constants ──────────────────────────────────────────────
    SUB_NAME_W  = 80   # width of member name column
    SUB_DAY_W   = 56   # width of each day column
    SUB_ROW_H   = 28   # height of each row
    SUB_HEADER_H = 36  # height of header row
    SUB_DUTY_SLOTS = ["日勤", "夜勤", "準夜"]

    def draw_sub_calendar(self):
        """サブカレンダー（メンバー×日付 グリッド）を描画する"""
        c = self.sub_canvas
        c.delete("all")
        year  = self.current_year
        month = self.current_month
        import calendar as _cal
        num_days = _cal.monthrange(year, month)[1]

        active = [m for m in self.members_data if m.get("active", True)]
        if not active:
            c.create_text(200, 100, text="メンバーがいません", font=("Meiryo UI", 12), fill="#888")
            return

        NW = self.SUB_NAME_W
        DW = self.SUB_DAY_W
        RH = self.SUB_ROW_H
        HH = self.SUB_HEADER_H
        total_w = NW + DW * num_days + 4
        total_h = HH + RH * len(active) + 4
        c.config(scrollregion=(0, 0, total_w, total_h))

        # ── header: day numbers ──
        c.create_rectangle(0, 0, NW, HH, fill="#2d6a4f", outline="")
        c.create_text(NW//2, HH//2, text=f"{month}月", fill="white",
                      font=("Meiryo UI", 10, "bold"))
        week_names = ["月","火","水","木","金","土","日"]
        holidays = self._current_holidays
        for d in range(1, num_days + 1):
            x0 = NW + (d-1)*DW
            x1 = x0 + DW
            wd = _cal.weekday(year, month, d)  # 0=Mon … 6=Sun
            is_sat = (wd == 5)
            is_sun_or_hol = (wd == 6) or (d in holidays)
            bg = "#e8f4fd" if is_sat else ("#fde8e8" if is_sun_or_hol else "#2d6a4f")
            fg = "#0055aa" if is_sat else ("#cc0000" if is_sun_or_hol else "white")
            c.create_rectangle(x0, 0, x1, HH, fill=bg, outline="#ccc")
            c.create_text((x0+x1)//2, HH//2 - 8, text=str(d),
                          fill=fg, font=("Meiryo UI", 9, "bold"))
            c.create_text((x0+x1)//2, HH//2 + 8, text=week_names[wd],
                          fill=fg, font=("Meiryo UI", 8))

        # ── rows: one per member ──
        for ri, member in enumerate(active):
            y0 = HH + ri * RH
            y1 = y0 + RH
            row_bg = "#f9fff9" if ri % 2 == 0 else "#f0f8f0"
            # name cell
            color = member.get("color", "#cccccc")
            c.create_rectangle(0, y0, NW, y1, fill=color, outline="#bbb")
            c.create_text(NW//2, (y0+y1)//2, text=member.get("display_name", member["name"]),
                          fill="black", font=("Meiryo UI", 9, "bold"))
            # day cells
            name = member.get("display_name", member["name"])
            for d in range(1, num_days + 1):
                x0 = NW + (d-1)*DW
                x1 = x0 + DW
                wd = _cal.weekday(year, month, d)
                is_sat = (wd == 5)
                is_sun_or_hol = (wd == 6) or (d in holidays)
                cell_bg = "#e8f4fd" if is_sat else ("#fde8e8" if is_sun_or_hol else row_bg)
                val = self.sub_cell_data.get((name, d, "duty"), "")
                if val:
                    cell_bg = member.get("color", "#cccccc")
                c.create_rectangle(x0, y0, x1, y1, fill=cell_bg, outline="#ddd",
                                   tags=(f"sub_cell_{name}_{d}",))
                if val:
                    c.create_text((x0+x1)//2, (y0+y1)//2, text=val,
                                  font=("Meiryo UI", 8), fill="black",
                                  tags=(f"sub_text_{name}_{d}",))

    def _sub_cell_at(self, cx, cy):
        """キャンバス座標 (cx,cy) からサブセルの (member_name, day) を返す。範囲外は None"""
        import calendar as _cal
        year  = self.current_year
        month = self.current_month
        num_days = _cal.monthrange(year, month)[1]
        active = [m for m in self.members_data if m.get("active", True)]
        NW = self.SUB_NAME_W
        DW = self.SUB_DAY_W
        RH = self.SUB_ROW_H
        HH = self.SUB_HEADER_H
        if cx < NW or cy < HH:
            return None
        col = (cx - NW) // DW   # 0-indexed day offset
        row = (cy - HH) // RH
        if col < 0 or col >= num_days or row < 0 or row >= len(active):
            return None
        day = col + 1
        member = active[row]
        name = member.get("display_name", member["name"])
        return (name, day)

    def _on_sub_canvas_click(self, event):
        cx = self.sub_canvas.canvasx(event.x)
        cy = self.sub_canvas.canvasy(event.y)
        result = self._sub_cell_at(cx, cy)
        if result is None:
            return
        name, day = result
        # toggle between duty values
        current = self.sub_cell_data.get((name, day, "duty"), "")
        duties = ["", "日勤", "夜勤", "準夜", "休み"]
        try:
            idx = duties.index(current)
        except ValueError:
            idx = 0
        next_val = duties[(idx + 1) % len(duties)]
        if next_val:
            self.sub_cell_data[(name, day, "duty")] = next_val
        else:
            self.sub_cell_data.pop((name, day, "duty"), None)
        self.draw_sub_calendar()

    def _on_sub_canvas_right_click(self, event):
        cx = self.sub_canvas.canvasx(event.x)
        cy = self.sub_canvas.canvasy(event.y)
        result = self._sub_cell_at(cx, cy)
        if result is None:
            return
        name, day = result
        menu = tk.Menu(self.root, tearoff=0)
        duties = ["日勤", "夜勤", "準夜", "休み", "有休", "研修", "外勤"]
        for d in duties:
            menu.add_command(label=d, command=lambda v=d, n=name, dy=day: self._set_sub_duty(n, dy, v))
        menu.add_separator()
        menu.add_command(label="クリア", command=lambda n=name, dy=day: self._set_sub_duty(n, dy, ""))
        menu.tk_popup(event.x_root, event.y_root)

    def _set_sub_duty(self, name: str, day: int, value: str):
        if value:
            self.sub_cell_data[(name, day, "duty")] = value
        else:
            self.sub_cell_data.pop((name, day, "duty"), None)
        self.draw_sub_calendar()

    def _on_sub_sync_btn(self):
        """メインカレンダー → サブカレンダーへ同期"""
        if messagebox.askyesno("同期確認", "メインカレンダーの内容でサブカレンダーを上書きします。\nよろしいですか？"):
            self._apply_main_to_sub()
            self.draw_sub_calendar()
            messagebox.showinfo("同期完了", "サブカレンダーを更新しました。")

    def _on_main_sync_btn(self):
        """サブカレンダー → メインカレンダーへ同期"""
        if messagebox.askyesno("同期確認", "サブカレンダーの内容でメインカレンダーを上書きします。\nよろしいですか？"):
            self._apply_sub_to_main()
            self._redraw_all_cells()
            self.draw_statistics()
            messagebox.showinfo("同期完了", "メインカレンダーを更新しました。")

    def _apply_main_to_sub(self):
        """メインの cell_data からサブ用データを生成"""
        self.sub_cell_data = {}
        duties_map = self._get_duties_from_main()
        for (name, day), duty in duties_map.items():
            self.sub_cell_data[(name, day, "duty")] = duty

    def _apply_sub_to_main(self):
        """サブの sub_cell_data をメインの cell_data に反映（LABELS[0] 行目）"""
        self._save_state()
        active_names = {m.get("display_name", m["name"]) for m in self.members_data if m.get("active", True)}
        for (name, day, slot), duty in self.sub_cell_data.items():
            if slot != "duty" or name not in active_names:
                continue
            # find cells that belong to this member on this day, clear them first
            keys_for_day = [(k, v) for k, v in self.cell_data.items()
                            if k[0] == day and v == name]
            for k, _ in keys_for_day:
                self.cell_data.pop(k, None)
                self.cell_colors.pop(k, None)
            # place in first label row, first col
            ck = (day, 0, 0)
            self.cell_data[ck] = name
            color = self.value_colors.get(name, self._get_or_assign_color(name))
            self.cell_colors[ck] = color

    def _get_duties_from_main(self) -> dict:
        """cell_data から {(name, day): duty_label} を構築する"""
        result = {}
        label_priority = {label: i for i, label in enumerate(self.LABELS)}
        for (day, label_idx, col), name in self.cell_data.items():
            if not name:
                continue
            cur = result.get((name, day))
            if cur is None:
                label = self.LABELS[label_idx] if label_idx < len(self.LABELS) else "日勤"
                result[(name, day)] = label
            else:
                # keep higher priority (lower index)
                cur_priority = label_priority.get(cur, 999)
                new_priority = label_priority.get(
                    self.LABELS[label_idx] if label_idx < len(self.LABELS) else "日勤", 999)
                if new_priority < cur_priority:
                    result[(name, day)] = self.LABELS[label_idx] if label_idx < len(self.LABELS) else "日勤"
        return result



# ═══════════════════════════════════════════════════════════════════════════
#  MemberManagementDialog
# ═══════════════════════════════════════════════════════════════════════════
class MemberManagementDialog:
    """メンバー管理ダイアログ（A/B/非常勤 チームパネル）"""

    def __init__(self, app: "CalendarApp"):
        self.app = app
        self.win = tk.Toplevel(app.root)
        self.win.title("👥 メンバー管理")
        self.win.geometry("760x580")
        self.win.resizable(True, True)
        self.win.configure(bg="#f5f7fa")
        self.win.transient(app.root)
        self.win.grab_set()
        self._build_ui()
        self._refresh()

    def _build_ui(self):
        top = tk.Frame(self.win, bg="#f5f7fa")
        top.pack(fill=tk.X, padx=12, pady=(10, 4))
        tk.Button(top, text="＋ メンバー追加", font=("Meiryo UI", 10, "bold"),
                  bg="#2d6a4f", fg="white", relief=tk.FLAT, padx=12, pady=4,
                  cursor="hand2", command=self._add_member).pack(side=tk.LEFT, padx=4)
        tk.Button(top, text="✏ 編集", font=("Meiryo UI", 10),
                  bg="#4a90d9", fg="white", relief=tk.FLAT, padx=12, pady=4,
                  cursor="hand2", command=self._edit_selected).pack(side=tk.LEFT, padx=4)
        tk.Button(top, text="🗑 削除", font=("Meiryo UI", 10),
                  bg="#c0392b", fg="white", relief=tk.FLAT, padx=12, pady=4,
                  cursor="hand2", command=self._delete_selected).pack(side=tk.LEFT, padx=4)
        tk.Button(top, text="↑ 上へ", font=("Meiryo UI", 10),
                  bg="#7f8c8d", fg="white", relief=tk.FLAT, padx=10, pady=4,
                  cursor="hand2", command=self._move_up).pack(side=tk.LEFT, padx=4)
        tk.Button(top, text="↓ 下へ", font=("Meiryo UI", 10),
                  bg="#7f8c8d", fg="white", relief=tk.FLAT, padx=10, pady=4,
                  cursor="hand2", command=self._move_down).pack(side=tk.LEFT, padx=4)

        notebook = ttk.Notebook(self.win)
        notebook.pack(fill=tk.BOTH, expand=True, padx=12, pady=6)
        self._team_lists: dict[str, tk.Listbox] = {}
        for team, label in [("A", "Aチーム"), ("B", "Bチーム"), ("非常勤", "非常勤")]:
            frame = tk.Frame(notebook, bg="white")
            notebook.add(frame, text=f"  {label}  ")
            sb = tk.Scrollbar(frame)
            sb.pack(side=tk.RIGHT, fill=tk.Y)
            lb = tk.Listbox(frame, font=("Meiryo UI", 11), selectmode=tk.SINGLE,
                            yscrollcommand=sb.set, activestyle="none",
                            selectbackground="#2d6a4f", selectforeground="white",
                            bg="white", relief=tk.FLAT, bd=0)
            lb.pack(fill=tk.BOTH, expand=True)
            sb.config(command=lb.yview)
            lb.bind("<Double-Button-1>", lambda e, t=team: self._edit_selected(team=t))
            self._team_lists[team] = lb
        self._notebook = notebook

        btm = tk.Frame(self.win, bg="#f5f7fa")
        btm.pack(fill=tk.X, padx=12, pady=(0, 10))
        tk.Button(btm, text="閉じる", font=("Meiryo UI", 10),
                  bg="#555", fg="white", relief=tk.FLAT, padx=16, pady=4,
                  cursor="hand2", command=self.win.destroy).pack(side=tk.RIGHT, padx=4)

    def _get_current_team(self) -> str:
        tab_idx = self._notebook.index(self._notebook.select())
        return ["A", "B", "非常勤"][tab_idx]

    def _refresh(self):
        for team, lb in self._team_lists.items():
            lb.delete(0, tk.END)
        for m in self.app.members_data:
            team = m.get("team", "A")
            lb = self._team_lists.get(team)
            if lb is None:
                continue
            status = "●" if m.get("active", True) else "○"
            lb.insert(tk.END, f"  {status}  {m.get('display_name', m['name'])}  [{m.get('skill_level','')}]")
            # color the entry
            idx = lb.size() - 1
            color = m.get("color", "#cccccc")
            try:
                lb.itemconfig(idx, bg=color)
            except Exception:
                pass

    def _selected_member(self, team: str = None) -> dict | None:
        t = team or self._get_current_team()
        lb = self._team_lists[t]
        sel = lb.curselection()
        if not sel:
            return None
        idx = sel[0]
        team_members = [m for m in self.app.members_data if m.get("team", "A") == t]
        if idx >= len(team_members):
            return None
        return team_members[idx]

    def _add_member(self):
        team = self._get_current_team()
        MemberEditDialog(self.app, self, team=team)

    def _edit_selected(self, event=None, team: str = None):
        member = self._selected_member(team)
        if member is None:
            messagebox.showwarning("選択なし", "編集するメンバーを選択してください。", parent=self.win)
            return
        MemberEditDialog(self.app, self, member=member)

    def _delete_selected(self):
        member = self._selected_member()
        if member is None:
            messagebox.showwarning("選択なし", "削除するメンバーを選択してください。", parent=self.win)
            return
        name = member.get("display_name", member["name"])
        if messagebox.askyesno("削除確認", f"「{name}」を削除しますか？\nカレンダーデータからも除去されます。", parent=self.win):
            self.app.members_data.remove(member)
            # remove from cell data
            keys = [k for k, v in self.app.cell_data.items() if v == name]
            for k in keys:
                self.app.cell_data.pop(k, None)
                self.app.cell_colors.pop(k, None)
            self.app.value_colors.pop(name, None)
            self.app._redraw_all_cells()
            self.app.draw_statistics()
            self._refresh()

    def _move_up(self):
        team = self._get_current_team()
        lb = self._team_lists[team]
        sel = lb.curselection()
        if not sel or sel[0] == 0:
            return
        idx = sel[0]
        team_members = [m for m in self.app.members_data if m.get("team", "A") == team]
        if idx >= len(team_members):
            return
        m = team_members[idx]
        gi = self.app.members_data.index(m)
        prev = team_members[idx - 1]
        gi_prev = self.app.members_data.index(prev)
        self.app.members_data[gi], self.app.members_data[gi_prev] = \
            self.app.members_data[gi_prev], self.app.members_data[gi]
        self._refresh()
        lb.selection_set(idx - 1)

    def _move_down(self):
        team = self._get_current_team()
        lb = self._team_lists[team]
        sel = lb.curselection()
        if not sel:
            return
        idx = sel[0]
        team_members = [m for m in self.app.members_data if m.get("team", "A") == team]
        if idx >= len(team_members) - 1:
            return
        m = team_members[idx]
        gi = self.app.members_data.index(m)
        nxt = team_members[idx + 1]
        gi_nxt = self.app.members_data.index(nxt)
        self.app.members_data[gi], self.app.members_data[gi_nxt] = \
            self.app.members_data[gi_nxt], self.app.members_data[gi]
        self._refresh()
        lb.selection_set(idx + 1)



# ═══════════════════════════════════════════════════════════════════════════
#  MemberEditDialog
# ═══════════════════════════════════════════════════════════════════════════
class MemberEditDialog:
    """メンバー追加／編集ダイアログ"""

    SKILL_OPTIONS = ["専門医", "後期研修医", "初期研修医", "看護師", "事務", "その他"]
    TEAM_OPTIONS  = ["A", "B", "非常勤"]

    def __init__(self, app: "CalendarApp", parent_dialog: "MemberManagementDialog",
                 member: dict = None, team: str = "A"):
        self.app = app
        self.parent_dialog = parent_dialog
        self.member = member  # None = new member
        self.win = tk.Toplevel(parent_dialog.win)
        self.win.title("✏ メンバー編集" if member else "＋ メンバー追加")
        self.win.geometry("420x420")
        self.win.resizable(False, False)
        self.win.configure(bg="#f5f7fa")
        self.win.transient(parent_dialog.win)
        self.win.grab_set()
        self._default_team = team
        self._color = member["color"] if member else "#aaccff"
        self._build_ui()

    def _build_ui(self):
        f = tk.Frame(self.win, bg="#f5f7fa", padx=20, pady=16)
        f.pack(fill=tk.BOTH, expand=True)

        def row(label, widget_factory, **kw):
            r = tk.Frame(f, bg="#f5f7fa")
            r.pack(fill=tk.X, pady=4)
            tk.Label(r, text=label, width=14, anchor="w",
                     font=("Meiryo UI", 10), bg="#f5f7fa").pack(side=tk.LEFT)
            w = widget_factory(r, **kw)
            w.pack(side=tk.LEFT, fill=tk.X, expand=True)
            return w

        self.name_var    = tk.StringVar(value=self.member["name"] if self.member else "")
        self.disp_var    = tk.StringVar(value=self.member.get("display_name","") if self.member else "")
        self.skill_var   = tk.StringVar(value=self.member.get("skill_level","専門医") if self.member else "専門医")
        self.team_var    = tk.StringVar(value=self.member.get("team","A") if self.member else self._default_team)
        self.request_var = tk.StringVar(value=self.member.get("personal_request","") if self.member else "")
        self.active_var  = tk.BooleanVar(value=self.member.get("active", True) if self.member else True)

        row("氏名 *", tk.Entry, textvariable=self.name_var, font=("Meiryo UI", 11))
        row("表示名", tk.Entry, textvariable=self.disp_var, font=("Meiryo UI", 11))
        row("職種",
            lambda p, **kw: ttk.Combobox(p, textvariable=self.skill_var,
                                          values=self.SKILL_OPTIONS, state="readonly",
                                          font=("Meiryo UI", 11)))
        row("チーム",
            lambda p, **kw: ttk.Combobox(p, textvariable=self.team_var,
                                          values=self.TEAM_OPTIONS, state="readonly",
                                          font=("Meiryo UI", 11)))
        row("リクエスト", tk.Entry, textvariable=self.request_var, font=("Meiryo UI", 11))

        # color picker row
        cr = tk.Frame(f, bg="#f5f7fa")
        cr.pack(fill=tk.X, pady=4)
        tk.Label(cr, text="カラー", width=14, anchor="w",
                 font=("Meiryo UI", 10), bg="#f5f7fa").pack(side=tk.LEFT)
        self.color_btn = tk.Button(cr, bg=self._color, width=6, relief=tk.GROOVE,
                                   cursor="hand2", command=self._pick_color)
        self.color_btn.pack(side=tk.LEFT, padx=(0, 8))
        tk.Label(cr, textvariable=tk.StringVar(value="← クリックで変更"),
                 font=("Meiryo UI", 9), fg="#888", bg="#f5f7fa").pack(side=tk.LEFT)

        # active checkbox
        ar = tk.Frame(f, bg="#f5f7fa")
        ar.pack(fill=tk.X, pady=4)
        tk.Label(ar, text="", width=14, bg="#f5f7fa").pack(side=tk.LEFT)
        tk.Checkbutton(ar, text="アクティブ（勤務表に表示）",
                       variable=self.active_var,
                       font=("Meiryo UI", 10), bg="#f5f7fa").pack(side=tk.LEFT)

        # buttons
        br = tk.Frame(f, bg="#f5f7fa")
        br.pack(fill=tk.X, pady=(16, 0))
        tk.Button(br, text="保存", font=("Meiryo UI", 10, "bold"),
                  bg="#2d6a4f", fg="white", relief=tk.FLAT, padx=16, pady=4,
                  cursor="hand2", command=self._save).pack(side=tk.RIGHT, padx=4)
        tk.Button(br, text="キャンセル", font=("Meiryo UI", 10),
                  bg="#888", fg="white", relief=tk.FLAT, padx=12, pady=4,
                  cursor="hand2", command=self.win.destroy).pack(side=tk.RIGHT, padx=4)

    def _pick_color(self):
        from tkinter import colorchooser
        color = colorchooser.askcolor(color=self._color, parent=self.win, title="カラーを選択")[1]
        if color:
            self._color = color
            self.color_btn.config(bg=color)

    def _save(self):
        name = self.name_var.get().strip()
        if not name:
            messagebox.showwarning("入力エラー", "氏名を入力してください。", parent=self.win)
            return
        display_name = self.disp_var.get().strip() or name

        if self.member is None:
            # new member
            new_id = str(self.app._next_member_id).zfill(3)
            self.app._next_member_id += 1
            m = {
                "id": new_id,
                "name": name,
                "display_name": display_name,
                "skill_level": self.skill_var.get(),
                "team": self.team_var.get(),
                "personal_request": self.request_var.get().strip(),
                "color": self._color,
                "active": self.active_var.get(),
            }
            self.app.members_data.append(m)
            self.app._apply_member_color_to_cells(display_name, self._color)
        else:
            old_display = self.member.get("display_name", self.member["name"])
            self.member["name"]             = name
            self.member["display_name"]     = display_name
            self.member["skill_level"]      = self.skill_var.get()
            self.member["team"]             = self.team_var.get()
            self.member["personal_request"] = self.request_var.get().strip()
            self.member["color"]            = self._color
            self.member["active"]           = self.active_var.get()
            self.app._apply_member_color_to_cells(display_name, self._color, old_display)

        self.parent_dialog._refresh()
        self.win.destroy()


# ═══════════════════════════════════════════════════════════════════════════
#  Entry point
# ═══════════════════════════════════════════════════════════════════════════
def main():
    root = tk.Tk()
    root.withdraw()   # hide until layout is ready

    # Splash / loading indication via title
    root.title("カレンダー管理システム — 起動中...")
    root.geometry("1400x860")
    root.minsize(900, 600)

    try:
        root.iconbitmap(default="calendar.ico")
    except Exception:
        pass

    app = CalendarApp(root)
    root.deiconify()
    root.mainloop()


if __name__ == "__main__":
    main()
